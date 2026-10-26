"""End-to-end check: a Semrush audit runs on just two exports (client backlinks + Backlink Gap matrix).

Builds synthetic exports with Semrush's column headers, runs prep.py and build.py with no --map and no
hand-made CSVs, and checks the workbook. The only extra inputs are the ones a person or Composio supplies:
config.json and the review decisions (decisions.json). Nothing is checked live.

Run from the repo root:  python3 -m unittest discover -s tests/backlink-audit -v
"""
import csv
import json
import subprocess
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

SKILL = Path(__file__).resolve().parents[2] / 'plugins/mavan-seo/skills/backlink-audit'
CLIENT = 'example-medspa.com'
COMPS = ['comp-one.com', 'comp-two.com', 'comp-three.com']
HOME = f'https://{CLIENT}/'

BACKLINK_HEADER = ['Page ascore', 'Source title', 'Source url', 'Target url', 'Anchor', 'External links',
                   'Internal links', 'Nofollow', 'Sponsored', 'Ugc', 'Text', 'Frame', 'Form', 'Image', 'Sitewide',
                   'First seen', 'Last seen', 'New link', 'Lost link']
GAP_HEADER = ['Domain', 'Domain ascore', CLIENT] + COMPS + ['Matches']

TARGETS = [HOME, f'{HOME}services/botox/', f'{HOME}services/']
N_REAL, N_SPAM = 25, 60


def link(src, target, anchor, title='', nofollow='false', lost='false'):
    return ['31', title, src, target, anchor, '3', '12', nofollow, 'false', 'false', 'true', 'false', 'false',
            'false', 'false', '2025-06-01', '2026-10-01', 'false', lost]


def backlinks_rows():
    rows = []
    for i in range(N_REAL):
        rows.append(link(f'https://localsite{i}.com/blog/post-{i}/', TARGETS[i % len(TARGETS)], 'Example Med Spa',
                         title=f'Local post {i}', nofollow='true' if i % 5 == 0 else 'false'))
    for i in range(N_SPAM):
        rows.append(link(f'https://spamnet{i}.xyz/dir/link-outreach-services-{10000 + i}/', HOME,
                         f'Premium Link Building Experts for Better Rankings {CLIENT}',
                         title='Premium Link Building Experts'))
    # lost links: a real site, a domain that still links from another page, and a spam one (left out)
    rows.append(link('https://lostsite.com/old/', f'{HOME}services/botox/', 'Example Med Spa', lost='true'))
    rows.append(link('https://lostsite.com/other/', HOME, 'Example Med Spa', lost='true'))
    rows.append(link('https://localsite1.com/old-post/', HOME, 'Example Med Spa', lost='true'))
    rows.append(link('https://spamnet1.xyz/dir/link-outreach-services-1/', HOME, 'Premium Link Building', lost='true'))
    return rows


def gap_rows():
    def r(dom, au, n):  # n competitors linking, Missing view: client column is always 0
        cells = [1 if i < n else 0 for i in range(3)]
        return [dom, au, 0] + cells + [n]
    rows = [r('tampabaylocalnews.com', 55, 3), r('floridachamber.org', 48, 2), r('bestmedspas-listing.com', 33, 2),
            r('frankfurtmedspa.com', 28, 2),             # contains "rank": must not be treated as junk
            r('smallbloglocal.com', 15, 1),              # 1 competitor, low authority: dropped unless decided
            r('onecomp-highauth.com', 62, 1)]
    for dom in ['cheap-backlinks.shop', 'seochecker.site', 'dapa-tools.online', 'rankbooster.xyz', 'bigdeals.store']:
        rows.append(r(dom, 12, 3))
    return rows


def write_csv(path, header, rows):
    with open(path, 'w', newline='', encoding='utf-8') as f:
        w = csv.writer(f)
        w.writerow(header)
        w.writerows(rows)


class TwoFileSemrushAudit(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory()
        cls.root = Path(cls.tmp.name)
        cls.work = cls.root / 'work'
        cls.work.mkdir()
        write_csv(cls.root / 'backlinks.csv', BACKLINK_HEADER, backlinks_rows())
        write_csv(cls.root / 'gap.csv', GAP_HEADER, gap_rows())
        (cls.work / 'config.json').write_text(json.dumps({
            'client': 'Example Med Spa', 'domain': CLIENT, 'tool': 'semrush', 'competitors': COMPS,
            'brand_terms': ['example med spa'], 'manual_action': 'none', 'manual_action_checked': '2026-10-26',
            'safe_domains': []}))
        cls.prep = cls.run_py('prep.py', '--config', str(cls.work / 'config.json'), '--out', str(cls.work),
                              '--backlinks', str(cls.root / 'backlinks.csv'), '--gap', str(cls.root / 'gap.csv'))
        (cls.work / 'decisions.json').write_text(json.dumps({
            'lost': {'lostsite.com': {'priority': 'High', 'do_this': 'Ask them to restore the link.'}},
            'opportunities': {
                # review step: every kept row of type Other needs a specific action
                'frankfurtmedspa.com': {'type': 'Industry publication / blog'},
                'onecomp-highauth.com': {'type': 'Industry publication / blog'},
                # added by hand: one competitor and low authority, so prep.py dropped it
                'smallbloglocal.com': {'type': 'Local news / media', 'do_this': 'Pitch a local story.',
                                       'priority': 'Medium'}}}))
        cls.out = cls.root / 'audit.xlsx'
        cls.build = cls.run_py('build.py', '--work', str(cls.work), '--out', str(cls.out))
        cls.wb = load_workbook(cls.out) if cls.out.exists() else None

    @classmethod
    def run_py(cls, script, *args):
        return subprocess.run([sys.executable, '-I', str(SKILL / 'scripts' / script), *args],
                              capture_output=True, text=True)

    @classmethod
    def tearDownClass(cls):
        cls.tmp.cleanup()

    def tab(self, name):
        ws = self.wb[name]
        rows = list(ws.iter_rows(values_only=True))
        return rows[0], [dict(zip(rows[0], r)) for r in rows[1:] if any(v is not None for v in r)]

    def test_prep_and_build_pass_with_two_exports(self):
        self.assertEqual(self.prep.returncode, 0, self.prep.stdout + self.prep.stderr)
        self.assertEqual(self.build.returncode, 0, self.build.stdout + self.build.stderr)
        self.assertIn('All checks passed', self.build.stdout)

    def test_no_missing_export_warning_for_derived_data(self):
        self.assertNotIn('MISSING EXPORTS', self.prep.stdout)

    def test_gap_authority_read_from_domain_ascore(self):
        _, rows = self.tab('Recommendations')
        opp = {r['Domain / Source']: r for r in rows if r['Type'] == 'Link opportunity'}
        self.assertEqual(opp['tampabaylocalnews.com']['AS'], 55)

    def test_no_live_check_files(self):
        self.assertFalse((self.work / 'status_needed.txt').exists())
        self.assertNotIn('Composio', self.prep.stdout + self.build.stdout)

    def test_lost_backlinks_from_export(self):
        _, rows = self.tab('Recommendations')
        lost = {r['Domain / Source']: r for r in rows if r['Type'] == 'Lost backlink'}
        self.assertEqual(set(lost), {'lostsite.com', 'localsite1.com'})  # the spam network is left out
        self.assertEqual(lost['lostsite.com']['Referring Domains or Competitors Linking'], 2)
        self.assertEqual(lost['lostsite.com']['Priority'], 'High')
        self.assertEqual(lost['lostsite.com']['Do This'], 'Ask them to restore the link.')
        self.assertIn('still links', lost['localsite1.com']['What We Found'])
        self.assertEqual(lost['localsite1.com']['Priority'], 'Low')
        self.assertIsNone(lost['lostsite.com']['Target URL'])
        summ = json.loads((self.work / 'prep_summary.json').read_text())
        self.assertEqual(summ['lost_links'], 4)

    def test_two_tabs_recommendations_first(self):
        self.assertEqual(self.wb.sheetnames, ['Recommendations', 'Disavow Review'])

    def test_template_columns(self):
        rec_header, _ = self.tab('Recommendations')
        self.assertEqual(list(rec_header), ['Priority', 'Type', 'Domain / Source', 'Client URL Affected', 'AS',
                                            'Referring Domains or Competitors Linking', 'What We Found', 'Do This',
                                            'Target URL', 'Evidence URL', 'Status'])
        dis_header, _ = self.tab('Disavow Review')
        self.assertEqual(list(dis_header), ['Referring Domain', 'Example Linking URL', 'Client URL Linked To',
                                            'Anchor Text (most common)', 'Links from Domain', 'AS',
                                            'Domain Organic Traffic', 'Dofollow?', 'First Seen', 'Tool Risk Score',
                                            'Tool Risk Markers', 'In GSC Links Report?', 'Why Flagged',
                                            'Pattern Group', 'Suggested Status', 'Final Decision', 'Reviewer Notes'])

    def test_baseline_counts_for_the_report(self):
        summ = json.loads((self.work / 'prep_summary.json').read_text())
        live_domains = N_REAL + N_SPAM  # the lost link is not counted
        self.assertEqual(summ['client_referring_domains'], live_domains)
        self.assertEqual(summ['client_backlinks'], live_domains)
        self.assertEqual(set(summ['gap_domains_by_competitor']), set(COMPS))
        self.assertTrue(summ['gap_domains_by_competitor']['comp-one.com'] > 0)

    def test_no_rollup_files_written(self):
        self.assertFalse((self.work / 'rollup.csv').exists())

    def test_spam_network_collapsed_into_noise_row(self):
        _, rows = self.tab('Disavow Review')
        self.assertFalse([r for r in rows if 'spamnet' in str(r['Referring Domain'])])
        noise = [r for r in rows if r['Pattern Group'] == 'Noise - auto-generated, ignore']
        self.assertEqual(len(noise), 1)
        self.assertGreaterEqual(noise[0]['Links from Domain'], N_SPAM)

    def test_junk_gap_domains_removed_but_real_ones_kept(self):
        _, rows = self.tab('Recommendations')
        opp = {r['Domain / Source'] for r in rows if r['Type'] == 'Link opportunity'}
        for junk in ['cheap-backlinks.shop', 'seochecker.site', 'dapa-tools.online', 'rankbooster.xyz',
                     'bigdeals.store']:
            self.assertNotIn(junk, opp)
        for good in ['tampabaylocalnews.com', 'floridachamber.org', 'frankfurtmedspa.com']:
            self.assertIn(good, opp)

    def test_decisions_can_add_a_one_competitor_domain(self):
        _, rows = self.tab('Recommendations')
        opp = {r['Domain / Source']: r for r in rows if r['Type'] == 'Link opportunity'}
        self.assertIn('smallbloglocal.com', opp)
        self.assertEqual(opp['smallbloglocal.com']['Priority'], 'Medium')

    def test_evidence_url_blank_on_semrush(self):
        _, rows = self.tab('Recommendations')
        for r in rows:
            if r['Type'] == 'Link opportunity':
                self.assertIsNone(r['Evidence URL'])

    def test_disavow_as_is_page_level_max_and_flagged(self):
        summ = json.loads((self.work / 'prep_summary.json').read_text())
        self.assertTrue(summ['authority_is_page_level'])
        self.assertIn('PAGE-level', self.prep.stdout + ' '.join(summ.get('notes', [])))
        _, rows = self.tab('Disavow Review')
        real = [r for r in rows if r['Pattern Group'] != 'Noise - auto-generated, ignore']
        for r in real:
            self.assertEqual(r['AS'], 31)

    def test_noise_row_has_no_as_and_counts_domains(self):
        _, rows = self.tab('Disavow Review')
        noise = [r for r in rows if r['Pattern Group'] == 'Noise - auto-generated, ignore'][0]
        self.assertIsNone(noise['AS'])
        self.assertEqual(noise['Links from Domain'], N_SPAM)

    def test_dofollow_yes_no_mixed(self):
        import pandas as pd
        sys.path.insert(0, str(SKILL / 'scripts'))
        from prep import dofollow_label
        self.assertEqual(dofollow_label(pd.DataFrame({'nofollow': [False, False]})), 'Yes')
        self.assertEqual(dofollow_label(pd.DataFrame({'nofollow': [True, True]})), 'No')
        self.assertEqual(dofollow_label(pd.DataFrame({'nofollow': [True, False]})), 'Mixed')

    def test_handoff_noise_patterns_and_tlds(self):
        sys.path.insert(0, str(SKILL / 'scripts'))
        from prep import NOISE_TEXT, junk_domain
        for s in ['example.com/all/356/x', 'Aged Domains for sale', 'SEO Domain Research', 'website stats for x',
                  'linkly.agency', 'Directory Pages Index', 'niche edits', 'tiered link building']:
            self.assertTrue(NOISE_TEXT.search(s), s)
        for d in ['x.icu', 'x.sbs', 'x.cfd', 'seotools.com']:
            self.assertTrue(junk_domain(d), d)
        for d in ['localdirectory.com', 'medspalistings.com', 'frank.com']:
            self.assertFalse(junk_domain(d), d)  # "directory"/"listing" are reviewed, not auto-removed


if __name__ == '__main__':
    unittest.main()
