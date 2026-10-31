#!/usr/bin/env python3
"""Build the single delivery workbook for a Content Prune & Merge run.

Usage:
  python make_xlsx.py --out "Client Content Pruning and Merging.xlsx" \
      --audit work/audit.tsv --recs work/recommendations.tsv \
      --existing work/existing_redirects.tsv [--not-in-crawl work/not_in_crawl.tsv]

Inputs are the tab-separated blocks the build scripts already write, with NO header row:
  audit.tsv              25 columns, URL ... Notes (references/rules.md order)
  recommendations.tsv    4 columns: Priority | Area | Problem | Do this
  existing_redirects.tsv 3 columns: Source URL | Target URL | Notes
  not_in_crawl.tsv       url | clicks | impressions | sessions | redirect target (a header row is allowed)

Tabs: Recommendations, Audit, Existing Redirects, Not in crawl. Headers match the template. The Audit
Area column and the Recommendations Pages / Where to find columns are written as plain values that
reproduce the template's formulas, so the file reads on its own. To use the template instead, copy
Audit columns B:Z into B2 of TEMPLATE - Audit (its Area formula fills column A) and the
Recommendations rows into B5.

Text that looks like a number or a formula (a leading = + - @) is stored as text, never evaluated.
"""
import argparse
import csv
import sys

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter

AUDIT_HEAD = ['Area', 'URL', 'Page Type', 'Page Type Reason', 'HTTP Status Code', 'Internal Linking URLs',
              'Internal Outgoing Links', 'Title', 'H1', 'URLs with Duplicate H1s', 'URLs with Duplicate Titles',
              'URLs with Similar Content', 'Conversions', 'New Users (GA4)', 'Sessions', 'Views (GA4)',
              'No. Keywords', 'Clicks', 'Impressions', 'Referring Domains', 'Suggested Next Step',
              'Suggested Target', 'Next Steps', 'Target URL', 'Reasoning', 'Notes']
# Positions in the 25-column block (without Area) that hold numbers
NUMERIC = {3, 4, 5, 8, 9, 10, 11, 12, 13, 14, 15, 16, 17, 18}
RECS_HEAD = ['#', 'Priority', 'Area', 'Problem', 'Do this', 'Pages', 'Where to find the pages', 'FINAL decision']
RECS_TITLE = 'Content Prune & Merge: Recommendations (TL;DR)'
EXIST_HEAD = ['Source URL', 'Target URL', 'Notes']
NIC_HEAD = ['URL', 'Clicks', 'Impressions', 'Sessions', 'Redirect target']
MERGE_REMOVE = ('Merge then Redirect', 'Remove then Redirect')


def read_tsv(path, width):
    rows = []
    with open(path, encoding='utf-8-sig', newline='') as f:
        for r in csv.reader(f, delimiter='\t', quoting=csv.QUOTE_NONE):
            if not any(c.strip() for c in r):
                continue
            rows.append((r + [''] * width)[:width])
    return rows


def area(row):
    """Same logic as the template's Area formula (Audit column A)."""
    d = row[21].strip() or row[19].strip()  # Next Steps, else Suggested Next Step
    n = row[24].lower()                      # Notes
    reason = row[2]                          # Page Type Reason
    r = d in MERGE_REMOVE
    out = []
    if d == 'Merge then Redirect':
        out.append('Merge pages')
    if d == 'Remove then Redirect':
        out.append('Remove pages')
    if d == 'Repurpose':
        out.append('Repurpose pages')
    if d == 'Update':
        out.append('Update pages')
    if d == 'Needs Discussion':
        out.append('Needs discussion')
    if 'blog outranks service' in n:
        out.append('Blog beats service')
    if 'city copy of' in n:
        out.append('City copies')
    if r and 'in main menu' in n:
        out.append('Menu cleanup')
    if r and 'already redirected to by' in n:
        out.append('Old redirects')
    if 'root-level' in reason.lower() or 'no rule matched' in reason.lower():
        out.append('Check page type')
    return ', '.join(out)


def put(ws, r, c, v, numeric=False, bold=False):
    cell = ws.cell(row=r, column=c)
    if numeric and v not in ('', None):
        try:
            x = float(str(v).replace(',', ''))
            cell.value = int(x) if x == int(x) else x
            return cell
        except ValueError:
            pass
    cell.value = v
    if isinstance(v, str):
        cell.data_type = 's'  # text, even when it starts with = + - @ or looks like a number
        cell.number_format = '@'
    if bold:
        cell.font = Font(bold=True)
    return cell


def header(ws, row, names):
    for i, h in enumerate(names, 1):
        c = ws.cell(row=row, column=i, value=h)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = ws.cell(row=row + 1, column=1)


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--out', required=True)
    ap.add_argument('--audit', required=True)
    ap.add_argument('--recs', required=True)
    ap.add_argument('--existing', required=True)
    ap.add_argument('--not-in-crawl')
    a = ap.parse_args()

    audit = read_tsv(a.audit, 25)
    recs = read_tsv(a.recs, 4)
    existing = read_tsv(a.existing, 3)
    areas = [area(r) for r in audit]

    wb = Workbook()

    ws = wb.active
    ws.title = 'Recommendations'
    ws.cell(row=1, column=1, value=RECS_TITLE).font = Font(bold=True, size=13)
    ws.cell(row=2, column=1, value='One row = one fix, highest priority first. To see the pages, go to the Audit tab and '
            'filter the Area column (column A) by the name in column C. Junk folders live on the Existing Redirects tab. '
            'Pick a FINAL decision for every row.')
    header(ws, 4, RECS_HEAD)
    junk_rules = sum(1 for r in existing if r[0].strip().endswith('*'))
    for i, (pri, ar, prob, todo) in enumerate(recs, 1):
        r = 4 + i
        put(ws, r, 1, i, numeric=True)
        for c, v in enumerate([pri, ar, prob, todo], 2):
            put(ws, r, c, v)
        if ar == 'Junk folders':
            n, where = junk_rules, 'Existing Redirects tab (Pages = wildcard rules)'
        else:
            n, where = sum(1 for x in areas if ar and ar.lower() in x.lower()), f'Audit tab > filter Area = {ar}'
        put(ws, r, 6, n, numeric=True)
        put(ws, r, 7, where)
        ws.cell(row=r, column=8).value = None
        for c in (4, 5):
            ws.cell(row=r, column=c).alignment = Alignment(wrap_text=True, vertical='top')
    widths(ws, [5, 10, 20, 60, 60, 8, 36, 24])

    ws = wb.create_sheet('Audit')
    header(ws, 1, AUDIT_HEAD)
    for i, (row, ar) in enumerate(zip(audit, areas), 2):
        put(ws, i, 1, ar)
        for j, v in enumerate(row):
            put(ws, i, j + 2, v, numeric=j in NUMERIC)
    widths(ws, [22, 50, 22, 30] + [12] * 3 + [34, 34] + [12] * 11 + [20, 40, 20, 40, 30, 70])
    ws.auto_filter.ref = ws.dimensions

    ws = wb.create_sheet('Existing Redirects')
    header(ws, 1, EXIST_HEAD)
    for i, row in enumerate(existing, 2):
        for j, v in enumerate(row, 1):
            put(ws, i, j, v)
    widths(ws, [60, 60, 90])

    ws = wb.create_sheet('Not in crawl')
    header(ws, 1, NIC_HEAD)
    if a.not_in_crawl:
        rows = read_tsv(a.not_in_crawl, 5)
        if rows and rows[0][0].strip().lower() in ('url', 'path'):
            rows = rows[1:]
        for i, row in enumerate(rows, 2):
            for j, v in enumerate(row, 1):
                put(ws, i, j, v, numeric=j in (2, 3, 4))
    widths(ws, [60, 10, 12, 10, 60])

    wb.save(a.out)
    print(f'Wrote {a.out}: {len(recs)} recommendations, {len(audit)} audit rows, {len(existing)} existing redirects')


if __name__ == '__main__':
    sys.exit(main())
