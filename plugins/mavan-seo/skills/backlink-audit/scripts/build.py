#!/usr/bin/env python3
"""Build the two-tab backlink audit workbook (Recommendations, Disavow Review) from prep.py output, status.tsv and decisions.json, then QA it.

Usage:
  python build.py --work work --out "CLIENT Backlink Audit Month YYYY.xlsx"

Broken backlinks need no export: every client URL the backlinks point at is checked live (status.tsv), and the ones
that answer 404, 410, 5xx, or 403 while the homepage answers 200, are listed as broken.

Prints "All checks passed" or a numbered list of problems. Fix each one in decisions.json (or by
re-running the status check) and run again; never hand-edit the workbook.
"""
import argparse
import math
import os
import sys

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import is_live, load_json, load_status, norm_url, reg_domain, bare_domain, save_json  # noqa: E402
from prep import OPP_TYPES, default_priority, dofollow_label, guess_type  # noqa: E402

PRIORITIES = ['High', 'Medium', 'Low']
STATUSES_DIS = ['Monitor', 'Review closer', 'Disavow candidate - matches manual action']
FINAL = ['Keep', 'Monitor', 'Request removal', 'Disavow']
REC_STATUS = ['Not started', 'In progress', 'Done', "Won't do"]
DEFAULT_DO = {
    'Local directory / citation': 'Claim or submit a listing; match the name, address and phone on the Google Business Profile.',
    'Manufacturer / provider locator': 'Ask the brand rep to add the client to the provider locator.',
    'Association / chamber / nonprofit': 'Join or get listed as a member; ask about sponsorship options.',
    'Local news / media': 'Pitch a story or expert quote; open the competitor link to see what they were featured for.',
    'Listicle / "best of" roundup': 'Pitch inclusion in the roundup; the listed competitors show the format.',
    'Sponsorship / event': 'Ask about sponsoring the event or organization.',
    'Podcast / interview': 'Pitch the practitioner as a guest.',
    'Industry publication / blog': 'Pitch an expert article or quote on a topic the client is known for.',
    'Education / .edu / .gov': 'Look for a resource page, scholarship or program the client can be part of.',
    'Partner / vendor': 'Ask the vendor to list the client as a customer or partner.',
}
HOME_TYPES = {'Local directory / citation', 'Manufacturer / provider locator', 'Association / chamber / nonprofit',
              'Sponsorship / event', 'Partner / vendor', 'Education / .edu / .gov', 'Podcast / interview'}
OPP_CAP = 75

HEAD_FILL = PatternFill('solid', fgColor='1F3864')
YELLOW = PatternFill('solid', fgColor='FFF2CC')
PFILL = {'High': 'F4CCCC', 'Medium': 'FFF2CC', 'Low': 'D9EAD3'}


def blank(v):
    return v is None or (isinstance(v, float) and math.isnan(v)) or (isinstance(v, str) and not v.strip())


def clean(v):
    if blank(v):
        return None
    if isinstance(v, float) and v.is_integer():
        return int(v)
    return v


def read(work, name):
    p = os.path.join(work, name)
    return pd.read_csv(p) if os.path.exists(p) and os.path.getsize(p) > 1 else pd.DataFrame()


def write_tab(ws, header, rows, widths, yellow=(), pct=(), nums=()):
    ws.append(header)
    for r in rows:
        ws.append([clean(r.get(h)) for h in header])
    for i, h in enumerate(header, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(name='Arial', bold=True, color='FFFFFF')
        c.fill = HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical='center')
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 14)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
        for c in row:
            h = header[c.column - 1]
            c.font = Font(name='Arial')
            c.alignment = Alignment(wrap_text=widths.get(h, 14) >= 30, vertical='top')
            if h in yellow:
                c.fill = YELLOW
            if h in pct:
                c.number_format = '0%'
            elif h in nums:
                c.number_format = '#,##0'
    ws.freeze_panes = 'A2'
    if rows:
        ws.auto_filter.ref = f'A1:{get_column_letter(len(header))}{len(rows) + 1}'


def dropdown(ws, header, col_name, values, n):
    if col_name not in header or n == 0:
        return
    letter = get_column_letter(header.index(col_name) + 1)
    dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f'{letter}2:{letter}{n + 1}')


BROKEN_COLS = ['url', 'code', 'refdomains', 'refdomains_total', 'best_domain', 'best_authority', 'evidence', 'map_target']
# Gap domains prep.py filtered out that decisions.json may still add back (a person's call, not a safety rule).
OVERRIDABLE = {'one competitor, low authority', 'junk domain (spam TLD or SEO-tool name)'}


def derive_broken(links, status, client, rmap, needed, noise=()):
    """Broken client URLs found from the live status of every URL the backlinks point at.

    Broken = 404, 410, 5xx, or 403 while the homepage answers 200 (checked after following redirects, so a
    redirect to a 404 counts). 429, errors, and 403 when the homepage is blocked too are bot protection or
    timeouts: they are reported as unchecked, not as broken. Returns (DataFrame like broken.csv, warnings).
    """
    if links.empty:
        return pd.DataFrame(columns=BROKEN_COLS), []
    home_url = norm_url(f'https://{client}/')
    home = status.get(home_url)
    if home is None:
        needed.add(home_url)
    walled = home is not None and home['status'] in ('403', '429', 'ERR', '')
    host = reg_domain(client)
    rows, unchecked, forbidden = [], [], []
    mine = links[links['target_url'].map(lambda u: reg_domain(u) == host)]
    for t, g in mine.groupby('target_url'):
        st = status.get(norm_url(t))
        if st is None:
            needed.add(norm_url(t))
            continue
        code = st['final_status'] or st['status']
        if code in ('', '429', 'ERR') or (code == '403' and walled):
            unchecked.append(t)
            continue
        if code == '403':
            forbidden.append(t)
        elif not (code in ('404', '410') or code.startswith('5')):
            continue
        real = g[~g['domain'].isin(noise)]  # priority counts real referring domains, not noise networks
        best = (real if not real.empty else g).sort_values('authority', ascending=False, na_position='last').iloc[0]
        rows.append({'url': t, 'code': code, 'refdomains': real['domain'].nunique(),
                     'refdomains_total': g['domain'].nunique(), 'best_domain': best['domain'],
                     'best_authority': best['authority'], 'evidence': best['source_url'],
                     'map_target': rmap.get(t, '')})
    warns = []
    if forbidden:
        warns.append(f'{len(forbidden)} client URL(s) answer 403 while the homepage answers 200, so they are listed '
                     f'as broken. Open a few in a browser to confirm, e.g. {", ".join(forbidden[:3])}.')
    if unchecked:
        warns.append(f'{len(unchecked)} client URL(s) could not be checked (403/429/error, likely bot protection) and '
                     f'are NOT listed as broken. Open them in a browser, e.g. {", ".join(unchecked[:3])}.')
    return pd.DataFrame(rows, columns=BROKEN_COLS), warns


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--work', required=True)
    ap.add_argument('--out', required=True)
    a = ap.parse_args()
    W = a.work
    cfg = load_json(os.path.join(W, 'config.json'))
    dec = load_json(os.path.join(W, 'decisions.json'), {}) or {}
    summ = load_json(os.path.join(W, 'prep_summary.json'), {})
    metric = summ.get('metric', 'DR')
    client = bare_domain(cfg['domain'])
    unnatural = cfg['manual_action'] == 'unnatural-to'
    status = load_status(W)
    problems, needed, soft = [], set(), []

    def live(u):
        if norm_url(u) not in status:
            needed.add(norm_url(u))
            return None
        return is_live(u, status)

    # ------------------------------------------------ Disavow Review
    dis = read(W, 'disavow.csv')
    noise = set(load_json(os.path.join(W, 'noise.json'), []) or []) | set(dec.get('noise', []))
    dd = dec.get('disavow', {})
    noise -= {d for d, o in dd.items() if not o.get('drop')}  # a domain added on review is not noise
    links = read(W, 'links.csv')
    dis_rows = []
    seen = set()
    for r in dis.to_dict('records'):
        d = r['Referring Domain']
        o = dd.get(d, {})
        if o.get('drop') or d in noise:
            if d in dec.get('noise', []):
                noise.add(d)
            continue
        for k, f in [('pattern_group', 'Pattern Group'), ('why_flagged', 'Why Flagged'), ('status', 'Suggested Status')]:
            if o.get(k):
                r[f] = o[k]
        seen.add(d)
        dis_rows.append(r)
    for d, o in dd.items():  # domains Claude added by hand
        if d in seen or o.get('drop'):
            continue
        g = links[links['domain'] == d] if not links.empty else pd.DataFrame()
        if g.empty:
            problems.append(f'decisions.json disavow adds {d}, which is not in the client backlinks export.')
            continue
        dis_rows.append({'Referring Domain': d, 'Example Linking URL': g['source_url'].iloc[0],
                         'Client URL Linked To': g['target_url'].mode().iloc[0],
                         'Anchor Text (most common)': g['anchor'].fillna('').mode().iloc[0],
                         'Links from Domain': len(g), metric: g['authority'].max(),
                         'Domain Organic Traffic': g['traffic'].max(), 'Dofollow?': dofollow_label(g),
                         'First Seen': '', 'Why Flagged': o.get('why_flagged', ''),
                         'Pattern Group': o.get('pattern_group', 'Added on review'),
                         'Suggested Status': o.get('status', 'Review closer')})
    if noise:
        dis_rows.append({'Referring Domain': f'{len(noise)} domain{"s" if len(noise) != 1 else ""} (see Why Flagged)', 'Links from Domain': len(noise),
                         'Why Flagged': 'Scraper, stats, "website worth" and SEO-tool pages. Google ignores these. '
                                        'Examples: ' + ', '.join(sorted(noise)[:8]),
                         'Pattern Group': 'Noise - auto-generated, ignore', 'Suggested Status': 'Monitor',
                         'Final Decision': 'Monitor'})
    order = {'Noise - auto-generated, ignore': 'zzz'}
    dis_rows.sort(key=lambda r: (order.get(r.get('Pattern Group'), r.get('Pattern Group') or ''),
                                 -(r.get('Links from Domain') or 0)))
    for r in dis_rows:
        s = r.get('Suggested Status')
        if s not in STATUSES_DIS:
            problems.append(f'Disavow Review {r["Referring Domain"]}: status "{s}" is not one of {STATUSES_DIS}.')
        if s == STATUSES_DIS[2] and not unnatural:
            problems.append(f'Disavow Review {r["Referring Domain"]}: marked as a disavow candidate, but the GSC '
                            f'Manual Actions result is "{cfg["manual_action"]}". Only Monitor or Review closer is allowed.')
    if unnatural and not any(r.get('Suggested Status') == STATUSES_DIS[2] for r in dis_rows):
        problems.append('Manual action is "Unnatural links to your site" but no row is a disavow candidate. '
                        'Mark the domains that match Google\'s examples in decisions.json.')

    # ------------------------------------------------ Recommendations: broken backlinks
    br = read(W, 'broken.csv')
    rmap = load_json(os.path.join(W, 'redirect_map.json'), {}) or {}
    derived_br, derived_warns = derive_broken(links, status, client, rmap, needed, noise)
    soft.extend(derived_warns)
    export_urls = set(br['url']) if 'url' in br else set()
    br = pd.concat([br, derived_br[~derived_br['url'].isin(export_urls)]], ignore_index=True) if not derived_br.empty else br
    all_broken = {norm_url(x) for x in br['url']} if 'url' in br else set()
    if 'url' in br:  # for the reviewer: everything build.py treats as a broken backlink, with who links to it
        br.to_csv(os.path.join(W, 'broken_found.csv'), index=False)
    bd = dec.get('broken', {})
    rec, fixed, broken_seen = [], [], set()
    for r in br.to_dict('records'):
        u = norm_url(r['url'])
        if u in broken_seen:
            problems.append(f'Broken URL listed twice: {u}')
        broken_seen.add(u)
        st = status.get(u)
        if st is None:
            needed.add(u)
            continue
        if st['status'] == '200' or st['final_status'] == '200':
            fixed.append(u + (f' (redirects to {st["final_url"]})' if st['final_url'] and st['final_url'] != u else ''))
            continue
        o = bd.get(u) or bd.get(r['url'], {})
        if o.get('action') == 'drop':
            continue
        tgt = norm_url(o.get('target') or (r['map_target'] if not blank(r.get('map_target')) else ''))
        if not tgt:
            problems.append(f'Broken URL {u} has no target: add it to decisions.json "broken" (who links to it: broken_found.csv).')
            continue
        ok = live(tgt)
        if ok is False:
            problems.append(f'Broken URL {u}: target {tgt} does not answer 200 without a redirect. Pick another target.')
        if tgt in all_broken:
            problems.append(f'Broken URL {u}: target {tgt} is itself a broken URL.')
        act = o.get('action', '301')
        src = r.get('best_domain') if not blank(r.get('best_domain')) else 'the linking site'
        if act == 'restore':
            do = f'Restore the page or 301 to {tgt}.'
        elif act == 'typo':
            do = f'301 the typo URL to {tgt}; optionally ask {src} to fix the link.'
        else:
            do = f'301 redirect to {tgt}.'
        if not blank(r.get('map_target')) and norm_url(r['map_target']) == tgt:
            do += ' (per prune & merge map)'
        if o.get('note'):
            do += ' ' + o['note']
        n = int(r['refdomains'] or 0)
        auth = r.get('best_authority')
        pr = o.get('priority') or ('High' if n >= 5 or (not blank(auth) and auth >= 40) else
                                   'Medium' if n >= 2 or (not blank(auth) and auth >= 20) else 'Low')
        code = '' if blank(r.get('code')) else str(r['code']).replace('.0', '')
        tot = 0 if blank(r.get('refdomains_total')) else int(r['refdomains_total'])
        total_note = f' ({tot} counting noise domains)' if tot > n else ''
        found = f'Returns {st["final_status"] or st["status"] or code or "an error"}; {n} referring domain{"s link" if n != 1 else " links"} here{total_note}'
        if not blank(r.get('best_domain')):
            found += f', strongest {r["best_domain"]}' + ('' if blank(auth) else f' ({metric} {auth:.0f})')
        rec.append({'Priority': pr, 'Type': 'Broken backlink', 'Domain / Source': clean(r.get('best_domain')),
                    'Client URL Affected': u, metric: auth, 'Referring Domains or Competitors Linking': n,
                    'What We Found': found + '.', 'Do This': do, 'Target URL': tgt,
                    'Evidence URL': clean(r.get('evidence'))})

    # ------------------------------------------------ Recommendations: link opportunities
    opp = read(W, 'opportunities.csv')
    od = dec.get('opportunities', {})
    dis_domains = {r['Referring Domain'] for r in dis_rows}
    home = f'https://{client}/'
    orows = []
    gap_by = {r['domain']: r for r in read(W, 'gap_all.csv').to_dict('records')}
    opp_records = opp.to_dict('records')
    have = {r['domain'] for r in opp_records}
    for d, o in od.items():  # domains added on review (e.g. one-competitor sites that prep.py dropped)
        if d in have or o.get('include') is False:
            continue
        g = gap_by.get(d)
        if g is None:
            problems.append(f'decisions.json opportunities adds {d}, which is not in the link gap export.')
            continue
        why = '' if blank(g['removed_reason']) else g['removed_reason']
        if why and why not in OVERRIDABLE:
            problems.append(f'decisions.json opportunities adds {d}, which was filtered out as "{why}" and can\'t be added.')
            continue
        t = o.get('type') or guess_type(d)
        n = int(g['n_competitors'] or 0)
        opp_records.append({'domain': d, 'authority': g['authority'], 'traffic': g['traffic'], 'n_competitors': n,
                            'competitors': '' if blank(g['competitors']) else g['competitors'], 'type_guess': t,
                            'priority': default_priority(n, g['authority'], t),
                            'evidence': '' if blank(g['evidence']) else g['evidence']})
    for r in opp_records:
        d = r['domain']
        o = od.get(d, {})
        if o.get('include') is False:
            continue
        t = o.get('type') or r['type_guess']
        if t not in OPP_TYPES:
            problems.append(f'Opportunity {d}: type "{t}" is not one of the listed types.')
        do = o.get('do_this') or DEFAULT_DO.get(t)
        if not do:
            problems.append(f'Opportunity {d} (type Other): write "do_this" in decisions.json, or set include false.')
            do = ''
        tgt = norm_url(o.get('target') or home)
        if t not in HOME_TYPES and not o.get('target') and t != 'Other':
            o_note = 'target defaults to the homepage; point editorial links at the matching service page'
            soft.append(f'Opportunity {d} ({t}): {o_note}.')
        if live(tgt) is False:
            problems.append(f'Opportunity {d}: target {tgt} does not answer 200 without a redirect.')
        if d in dis_domains:
            problems.append(f'Opportunity {d} is also on the Disavow Review tab.')
        comps = r['competitors'] if not blank(r['competitors']) else ''
        n = int(r['n_competitors'] or 0)
        found = f'{t}. Links to {comps or str(n) + " competitors"} but not the client.'
        if o.get('found'):
            found += ' ' + o['found']
        orows.append({'Priority': o.get('priority') or r['priority'], 'Type': 'Link opportunity', 'Domain / Source': d,
                      'Client URL Affected': None, metric: r['authority'],
                      'Referring Domains or Competitors Linking': n, 'What We Found': found, 'Do This': do,
                      'Target URL': tgt, 'Evidence URL': o.get('evidence') or clean(r.get('evidence'))})
    porder = {'High': 0, 'Medium': 1, 'Low': 2}
    orows.sort(key=lambda x: (porder.get(x['Priority'], 9), -(x['Referring Domains or Competitors Linking'] or 0),
                              -(x[metric] if not blank(x[metric]) else -1)))
    extra = max(0, len(orows) - OPP_CAP)
    orows = orows[:OPP_CAP]
    rec += orows
    for x in rec:
        if x['Priority'] not in PRIORITIES:
            problems.append(f'{x["Domain / Source"] or x["Client URL Affected"]}: priority "{x["Priority"]}" invalid.')
    rec.sort(key=lambda x: (porder.get(x['Priority'], 9), 0 if x['Type'] == 'Broken backlink' else 1,
                            -(x[metric] if not blank(x[metric]) else -1)))

    if needed:
        p = os.path.join(W, 'status_needed.txt')
        with open(p, 'w') as f:
            f.write('\n'.join(sorted(needed)) + '\n')
        problems.insert(0, f'{len(needed)} URLs have no live status yet. Check them in the Composio workbench, add '
                           f'the lines to {W}/status.tsv and run build.py again. They are listed in {p}.')

    # ------------------------------------------------ workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Recommendations'
    ch = ['Priority', 'Type', 'Domain / Source', 'Client URL Affected', metric, 'Referring Domains or Competitors Linking',
          'What We Found', 'Do This', 'Target URL', 'Evidence URL', 'Status']
    write_tab(ws, ch, rec, {'Type': 16, 'Domain / Source': 26, 'Client URL Affected': 36, 'What We Found': 50,
                            'Do This': 50, 'Target URL': 34, 'Evidence URL': 36, 'Status': 14,
                            'Referring Domains or Competitors Linking': 16},
              yellow={'Status'}, nums={'Referring Domains or Competitors Linking'})
    for rr in range(2, len(rec) + 2):
        p = ws.cell(row=rr, column=1)
        if p.value in PFILL:
            p.fill = PatternFill('solid', fgColor=PFILL[p.value])
    dropdown(ws, ch, 'Status', REC_STATUS, len(rec))
    dropdown(ws, ch, 'Priority', PRIORITIES, len(rec))

    ws = wb.create_sheet('Disavow Review')
    dh = ['Referring Domain', 'Example Linking URL', 'Client URL Linked To', 'Anchor Text (most common)',
          'Links from Domain', metric, 'Domain Organic Traffic', 'Dofollow?', 'First Seen', 'Tool Risk Score',
          'Tool Risk Markers', 'In GSC Links Report?', 'Why Flagged', 'Pattern Group', 'Suggested Status',
          'Final Decision', 'Reviewer Notes']
    write_tab(ws, dh, dis_rows, {'Referring Domain': 26, 'Example Linking URL': 40, 'Client URL Linked To': 34,
                                 'Anchor Text (most common)': 26, 'Why Flagged': 50, 'Pattern Group': 28,
                                 'Suggested Status': 22, 'Final Decision': 16, 'Reviewer Notes': 34,
                                 'Tool Risk Markers': 30},
              yellow={'Final Decision', 'Reviewer Notes'}, nums={'Links from Domain', 'Domain Organic Traffic'})
    dropdown(ws, dh, 'Final Decision', FINAL, len(dis_rows))
    wb.save(a.out)

    # ------------------------------------------------ report
    counts = {'disavow_review': len(dis_rows),
              'broken_backlinks': sum(1 for x in rec if x['Type'] == 'Broken backlink'),
              'link_opportunities': len(orows), 'opportunities_over_cap': extra, 'already_fixed': fixed}
    save_json(os.path.join(W, 'build_summary.json'), counts)
    print(f'Saved {a.out}')
    for k, v in counts.items():
        print(f'  {k}: {len(v) if isinstance(v, list) else v}')
    if fixed:
        print('  Already fixed (dropped):', '; '.join(fixed[:20]))
    if soft:
        print(f'\n{len(soft)} warnings (not blocking; review them):')
        for w in soft[:15]:
            print('  -', w)
        if len(soft) > 15:
            print(f'  ... and {len(soft) - 15} more of the same kind')
    if problems:
        print(f'\n{len(problems)} problems:')
        for i, p in enumerate(problems, 1):
            print(f'  {i}. {p}')
        sys.exit(1)
    print('\nAll checks passed')


if __name__ == '__main__':
    main()
