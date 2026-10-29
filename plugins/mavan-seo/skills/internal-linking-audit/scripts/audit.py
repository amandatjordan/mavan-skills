#!/usr/bin/env python3
"""Mavan internal linking audit, crawl data only (no page reading).

  python audit.py prep --urls URLS.csv --links LINKS.csv --base https://example.com [--body-regex REGEX]
      Normalises Sitebulb or Screaming Frog exports into work/. Without --body-regex it prints the
      link-path patterns so you can pick the one that means "inside the main content".
  python audit.py candidates --base https://example.com --priority priority.csv --service-regex REGEX
      [--gallery-regex REGEX] [--low-value-regex REGEX] [--exclude-regex REGEX]
      [--gsc-pages gsc_pages.csv] [--gsc-queries gsc_queries.csv]
      Writes work/candidates.csv, work/service_links.csv, work/gallery.csv, work/wrong_target.csv,
      work/baseline.csv for Claude to review. Every URL pattern comes from the site's own URLs
      (see the folder table prep prints); nothing about the site's structure is assumed.
      --gallery-regex    gallery / portfolio / case study / review pages (skip = no gallery check)
      --low-value-regex  pages a service-named anchor should not point to, e.g. blog, gallery,
                         reviews, contact (skip = no wrong-target check)
      --exclude-regex    pages that are never worth editing, e.g. tag/category/pagination archives
  python audit.py build --base https://example.com --rows work/rows.csv --client NAME --out FILE.xlsx
      [--findings work/findings.txt]
      Checks every row (URLs exist, no existing body link, High/Low) and builds the sheet.

priority.csv: url[,terms]  terms = '|'-separated phrases for the service (optional; derived from the H1).
gsc_pages.csv: page,clicks (top pages by clicks)   gsc_queries.csv: page,query,clicks,impressions (priority topics only)
"""
import argparse, os, re, sys, warnings
warnings.filterwarnings("ignore", message="This pattern")
import pandas as pd

W = 'work'
ACTIONS = ['Add link', 'Add link (template)', 'Change link target']
TYPES = ['Related content', 'Related service', 'Gallery & reviews', 'Wrong target']

def norm(u, base):
    if not isinstance(u, str) or not u.strip(): return ''
    u = u.strip()
    if u.startswith('/'): u = base + u
    u = re.sub(r'^http://', 'https://', u); u = re.sub(r'[#?].*$', '', u)
    m = re.match(r'(https://[^/]+)(.*)', u)
    return (m.group(1).lower() + m.group(2) if m else u).rstrip('/')

def pick(df, *names):
    for n in names:
        if n in df.columns: return df[n]
    return pd.Series([None] * len(df), index=df.index)

def terms_rx(terms):
    return re.compile(r'\b(?:' + '|'.join(re.escape(t).replace(r'\ ', r'[\s-]+') for t in terms if t) + r')s?\b', re.I)

def derive_terms(h1, url):
    t = re.split(r'\s+in\s+|\s*[|,–-]\s*', str(h1), maxsplit=1)[0].strip().lower() if h1 and str(h1) != 'nan' else ''
    t = re.sub(r'\s+(services?|treatments?|procedures?|surgery|options?|solutions?)$', '', t)
    if len(t) < 3:
        t = re.sub(r'-(in-)?[a-z]+-[a-z]{2}$', '', url.rstrip('/').split('/')[-1]).replace('-', ' ')
    return [t]

def prep(a):
    os.makedirs(W, exist_ok=True); base = a.base.rstrip('/')
    u = pd.read_csv(a.urls, encoding='utf-8-sig', low_memory=False)
    U = pd.DataFrame({'url': pick(u, 'Address', 'URL').map(lambda x: norm(x, base)),
                      'status': pd.to_numeric(pick(u, 'Status Code', 'HTTP Status Code'), errors='coerce'),
                      'indexable': pick(u, 'Indexability', 'Indexable Status').astype(str).str.lower().eq('indexable'),
                      'title': pick(u, 'Title 1', 'Title').astype(str), 'h1': pick(u, 'H1-1', 'H1').astype(str),
                      'words': pd.to_numeric(pick(u, 'Word Count', 'No. Content Words'), errors='coerce')})
    l = pd.read_csv(a.links, encoding='utf-8-sig', low_memory=False)
    if 'Linking URL' in l.columns:
        L = pd.DataFrame({'src': l['Linking URL'], 'dst': l['Target URL'], 'anchor': l['Anchor Text'],
                          'position': pick(l, 'Location'), 'path': pick(l, 'CSS Path')})
    elif 'Source' in l.columns:
        if 'Type' in l.columns: l = l[l['Type'].astype(str).str.lower().eq('hyperlink')]
        L = pd.DataFrame({'src': l['Source'], 'dst': l['Destination'], 'anchor': pick(l, 'Anchor'),
                          'position': pick(l, 'Link Position'), 'path': pick(l, 'Link Path')})
    else:
        sys.exit(f'Unrecognised links export: {list(l.columns)[:12]}')
    L['src'] = L.src.map(lambda x: norm(x, base)); L['dst'] = L.dst.map(lambda x: norm(x, base))
    L = L[L.dst.str.startswith(base)].copy()
    for c in ('anchor', 'position', 'path'): L[c] = L[c].fillna('').astype(str).str.strip()
    L['pattern'] = L.path.str.replace(r':nth-child\(\d+\)|\[\d+\]', '', regex=True).str.replace(r'\d+', 'N', regex=True)
    L['body'] = L.path.str.contains(a.body_regex, regex=True) if a.body_regex else L.position.str.lower().eq('content')
    U.to_pickle(f'{W}/urls.pkl'); L.to_pickle(f'{W}/links.pkl')
    print(f'urls={len(U)} (200 + indexable: {int(((U.status == 200) & U.indexable).sum())}), internal links={len(L)}')
    print('link positions:', L.position.value_counts().to_dict())
    c = L[L.position.str.lower().eq('content')]
    if len(c):
        print('\n"Content" link patterns (pick --body-regex that keeps only links inside the article/body):')
        g = c.groupby('pattern').agg(links=('dst', 'size'), pages=('src', 'nunique'), example=('anchor', lambda s: s.iloc[0][:50]))
        print(g.sort_values('links', ascending=False).head(25).to_string())
    if a.body_regex:
        b = L[L.body]
        print(f'\nbody links: {len(b)} from {b.src.nunique()} pages; by section:',
              b.src.str.replace(base, '').str.split('/').str[1].value_counts().head(8).to_dict())
    live = U[U.status == 200]
    f = live.url.str.replace(base, '').str.split('/').str[1].replace('', '(home)')
    print('\nlive pages by top folder (use these to write the service, gallery, low-value and exclude regexes):')
    print(pd.DataFrame({'folder': f, 'url': live.url}).groupby('folder').agg(pages=('url', 'size'), example=('url', lambda s: s.iloc[0].replace(base, '')))
          .sort_values('pages', ascending=False).head(30).to_string())

def candidates(a):
    base = a.base.rstrip('/')
    U = pd.read_pickle(f'{W}/urls.pkl'); L = pd.read_pickle(f'{W}/links.pkl')
    live = U[(U.status == 200) & U.indexable]
    svc_re = re.compile(a.service_regex)
    services = sorted(u for u in live.url if svc_re.search(u.replace(base, '')))
    pr = pd.read_csv(a.priority).fillna('')
    if 'terms' not in pr.columns: pr['terms'] = ''
    pr['url'] = pr.url.map(lambda x: norm(x, base))
    h1 = dict(zip(U.url, U.h1)); title = dict(zip(U.url, U.title))
    terms = {s: derive_terms(h1.get(s), s) for s in services}
    for _, r in pr.iterrows():
        terms[r.url] = [t.strip().lower() for t in str(r.terms).split('|') if t.strip()] or derive_terms(h1.get(r.url), r.url)
    targets = list(pr.url)
    body = L[L.body]
    body_pairs = set(zip(body.src, body.dst))
    clicks = {}
    if a.gsc_pages:
        g = pd.read_csv(a.gsc_pages); g['page'] = g.page.map(lambda x: norm(x, base)); clicks = g.groupby('page').clicks.sum().to_dict()
    gq = None
    if a.gsc_queries:
        gq = pd.read_csv(a.gsc_queries); gq['page'] = gq.page.map(lambda x: norm(x, base))
    def match(df, rx): return df.url.str.replace(base, '').str.contains(rx, regex=True) if rx else pd.Series(False, index=df.index)
    content = live[~match(live, a.gallery_regex) & ~match(live, a.exclude_regex)]

    rows = []
    for t in targets:
        rx = terms_rx(terms[t])
        for _, p in content.iterrows():
            s = p.url
            if s == t or (s, t) in body_pairs: continue
            why = []
            if rx.search(' '.join([s.replace(base, '').replace('-', ' '), p.title, p.h1])): why.append('title/H1/URL')
            if gq is not None:
                q = gq[(gq.page == s) & gq['query'].str.contains(rx)]
                if len(q): why.append(f'ranks for {len(q)} related searches, e.g. "{q.sort_values("impressions", ascending=False)["query"].iloc[0]}"')
            if why:
                rows.append({'type': 'Related service' if s in services else 'Related content', 'target': t, 'source': s,
                             'source_title': p.title, 'source_clicks': clicks.get(s, 0) if a.gsc_pages else '', 'matched_on': '; '.join(why)})
    C = pd.DataFrame(rows)
    if len(C): C = C.sort_values(['target', 'type', 'source_clicks'], ascending=[True, True, False])
    C.to_csv(f'{W}/candidates.csv', index=False)

    SL = [{'priority_page': t, 'service': s, 'service_h1': h1.get(s, ''), 'priority_links_to_it': (t, s) in body_pairs,
           'it_links_to_priority': (s, t) in body_pairs} for t in targets for s in services if s != t]
    pd.DataFrame(SL).to_csv(f'{W}/service_links.csv', index=False)

    G = []
    gal = U[(U.status == 200) & match(U, a.gallery_regex) & ~U.url.isin(services)] if a.gallery_regex else U.iloc[0:0]
    for t in targets if a.gallery_regex else []:
        rx = terms_rx(terms[t])
        rel = gal[[bool(rx.search(' '.join([u.replace(base, '').replace('-', ' '), ti, h]))) for u, ti, h in zip(gal.url, gal.title, gal.h1)]]
        up = L[L.src.isin(rel.url) & (L.dst == t) & ~L.position.str.lower().isin(['header', 'navigation', 'nav', 'footer'])]
        G.append({'priority_page': t, 'gallery_review_pages': len(rel), 'linking_up_outside_menu': up.src.nunique(),
                  'examples': ' | '.join(rel.url.head(3).str.replace(base, ''))})
    pd.DataFrame(G).to_csv(f'{W}/gallery.csv', index=False)

    WT = []
    for _, r in (body.iterrows() if a.low_value_regex else []):
        if r.dst in services or not re.search(a.low_value_regex, r.dst.replace(base, '')): continue
        for s in services:
            if terms_rx(terms[s]).search(r.anchor):
                WT.append({'source': r.src, 'anchor': r.anchor, 'current_target': r.dst, 'suggested_target': s,
                           'source_already_links_to_suggested': (r.src, s) in body_pairs, 'is_priority': s in targets}); break
    pd.DataFrame(WT).to_csv(f'{W}/wrong_target.csv', index=False)

    B = [{'priority_page': t, 'body_links_in': int((body.dst == t).sum()), 'pages_linking_in_body': body[body.dst == t].src.nunique(),
          'body_links_out': int((body.src == t).sum()),
          'topic_clicks_3mo': int(gq[gq.page == t].clicks.sum()) if gq is not None else ''} for t in targets]
    pd.DataFrame(B).to_csv(f'{W}/baseline.csv', index=False)

    print(f'services={len(services)} priority pages={len(targets)} content pages scanned={len(content)}')
    print(pd.DataFrame(B).to_string(index=False))
    if len(C): print('\ncandidates by page/type:', C.groupby(['target', 'type']).size().to_dict())
    if not a.gallery_regex: print('gallery check skipped (no --gallery-regex)')
    if not a.low_value_regex: print('wrong-target check skipped (no --low-value-regex)')
    print('gallery/review link-up:', [(g['priority_page'].replace(base, ''), g['gallery_review_pages'], g['linking_up_outside_menu']) for g in G])
    print(f'wrong-target anchors: {len(WT)} (source already links to the service: {sum(w["source_already_links_to_suggested"] for w in WT)})')
    print('\nterms used:'); [print(' ', t.replace(base, ''), '->', terms[t]) for t in targets]

def build(a):
    base = a.base.rstrip('/')
    U = pd.read_pickle(f'{W}/urls.pkl'); L = pd.read_pickle(f'{W}/links.pkl')
    ok = set(U[U.status == 200].url); body_pairs = set(zip(L[L.body].src, L[L.body].dst))
    R = pd.read_csv(a.rows).fillna('')
    need = ['priority', 'type', 'priority_page', 'action', 'source', 'target', 'why']
    miss = [c for c in need if c not in R.columns]
    if miss: sys.exit(f'rows.csv is missing columns: {miss}')
    R['source'] = R.source.map(lambda x: norm(x, base)); R['target'] = R.target.map(lambda x: norm(x, base))
    errs = []
    for i, r in R.iterrows():
        tag = f'row {i + 2} ({r.source.replace(base, "") or "/"} -> {r.target.replace(base, "")})'
        if r.priority not in ('High', 'Low'): errs.append(f'{tag}: priority must be High or Low')
        if r.action not in ACTIONS: errs.append(f'{tag}: action must be one of {ACTIONS}')
        if r.type not in TYPES: errs.append(f'{tag}: type must be one of {TYPES}')
        if r.source not in ok: errs.append(f'{tag}: page to edit is not a live 200 page in the crawl')
        if r.target not in ok: errs.append(f'{tag}: link target is not a live 200 page in the crawl')
        if r.action != 'Add link (template)' and (r.source, r.target) in body_pairs:
            errs.append(f'{tag}: page already links to the target in its body copy')
    dup = R[R.duplicated(['source', 'target'], keep=False)]
    for _, r in dup.iterrows(): errs.append(f'duplicate row: {r.source} -> {r.target}')
    for e in errs: print('ERROR', e)
    print(f'{len(R)} rows checked, {len(errs)} errors')
    if errs: sys.exit(1)

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.formatting.rule import CellIsRule
    R = R.assign(o=R.priority.map({'High': 0, 'Low': 1})).sort_values(['o', 'priority_page', 'type', 'source']).drop(columns='o')
    F = 'Arial'; H = Font(name=F, bold=True, color='FFFFFF'); HF = PatternFill('solid', fgColor='1F3A5F')
    N = Font(name=F, size=10); Wr = Alignment(wrap_text=True, vertical='top'); line = Border(bottom=Side(style='thin', color='D0D0D0'))
    def head(ws, cols, widths):
        for i, (h, w) in enumerate(zip(cols, widths), 1):
            c = ws.cell(1, i, h); c.font = H; c.fill = HF; c.alignment = Alignment(wrap_text=True, vertical='center')
            ws.column_dimensions[c.column_letter].width = w
    wb = Workbook(); s = wb.active; s.title = 'Summary'
    s['A1'] = f'Internal Linking Audit — {a.client}'; s['A1'].font = Font(name=F, bold=True, size=14)
    for i, h in enumerate(['Priority page', 'High', 'Low', 'Total', 'Done'], 1):
        c = s.cell(3, i, h); c.font = H; c.fill = HF
    pages = list(dict.fromkeys(R.priority_page))
    for j, p in enumerate(pages):
        r = 4 + j; s.cell(r, 1, p).font = N
        s.cell(r, 2, f'=COUNTIFS(Implementation!$C:$C,$A{r},Implementation!$A:$A,"High")').font = N
        s.cell(r, 3, f'=COUNTIFS(Implementation!$C:$C,$A{r},Implementation!$A:$A,"Low")').font = N
        s.cell(r, 4, f'=B{r}+C{r}').font = N
        s.cell(r, 5, f'=COUNTIFS(Implementation!$C:$C,$A{r},Implementation!$J:$J,"Done")').font = N
    t = 4 + len(pages); s.cell(t, 1, 'Total').font = Font(name=F, bold=True)
    for k, col in enumerate('BCDE', 2): s.cell(t, k, f'=SUM({col}4:{col}{t - 1})').font = Font(name=F, bold=True)
    r0 = t + 2; s.cell(r0, 1, 'Key findings').font = Font(name=F, bold=True)
    fl = [x.strip() for x in open(a.findings).read().splitlines() if x.strip()] if a.findings else []
    for i, f in enumerate(fl):
        c = s.cell(r0 + 1 + i, 1, f'{i + 1}. {f}'); c.font = N; c.alignment = Wr
        s.merge_cells(start_row=r0 + 1 + i, start_column=1, end_row=r0 + 1 + i, end_column=5); s.row_dimensions[r0 + 1 + i].height = 30
    hu = r0 + 2 + len(fl)
    s.cell(hu, 1, 'How to use: open each "Page to edit", find a natural spot in the copy, and link to "Link to" with descriptive anchor text. '
                  'If the page already links there, mark the row Rejected. Set Status as you go; the counts update automatically.').font = Font(name=F, size=10, italic=True)
    s.merge_cells(start_row=hu, start_column=1, end_row=hu, end_column=5); s.cell(hu, 1).alignment = Wr; s.row_dimensions[hu].height = 45
    s.column_dimensions['A'].width = 34
    for col in 'BCDE': s.column_dimensions[col].width = 10
    im = wb.create_sheet('Implementation')
    head(im, ['Priority', 'Type', 'Priority page', 'Action', 'Page to edit', 'Link to', 'Why', 'Owner', 'Date done', 'Status'],
         [9, 16, 20, 18, 55, 55, 60, 12, 11, 11])
    for i, r in enumerate(R.itertuples(), 2):
        for j, v in enumerate([r.priority, r.type, r.priority_page, r.action, r.source, r.target, r.why, '', '', 'To do'], 1):
            c = im.cell(i, j, v); c.font = N; c.alignment = Wr; c.border = line
    n = len(R) + 1
    dv = DataValidation(type='list', formula1='"To do,In progress,Done,Rejected"', allow_blank=True); im.add_data_validation(dv); dv.add(f'J2:J{n}')
    im.conditional_formatting.add(f'A2:A{n}', CellIsRule(operator='equal', formula=['"High"'], fill=PatternFill('solid', fgColor='F8D7DA')))
    im.conditional_formatting.add(f'A2:A{n}', CellIsRule(operator='equal', formula=['"Low"'], fill=PatternFill('solid', fgColor='E2E3E5')))
    im.conditional_formatting.add(f'J2:J{n}', CellIsRule(operator='equal', formula=['"Done"'], fill=PatternFill('solid', fgColor='D4EDDA')))
    im.freeze_panes = 'E2'; im.auto_filter.ref = f'A1:J{n}'
    b = wb.create_sheet('Baseline')
    bl = pd.read_csv(f'{W}/baseline.csv').fillna('')
    head(b, ['Priority page', 'Body links in', 'Pages linking in body copy', 'Body links out', 'Clicks from topic searches (last 3 months)'], [62, 14, 16, 14, 20])
    for i, r in enumerate(bl.itertuples(), 2):
        for j, v in enumerate([r.priority_page, r.body_links_in, r.pages_linking_in_body, r.body_links_out, r.topic_clicks_3mo], 1): b.cell(i, j, v).font = N
    b.cell(len(bl) + 3, 1, 'From the crawl export. Sitebulb can hide a body link that uses the same anchor as a menu link, so these may undercount.').font = Font(name=F, size=9, italic=True)
    if a.note: b.cell(len(bl) + 4, 1, a.note).font = Font(name=F, size=9, italic=True)
    wb.save(a.out); print(f'saved {a.out}')

if __name__ == '__main__':
    ap = argparse.ArgumentParser(); sub = ap.add_subparsers(dest='cmd', required=True)
    p = sub.add_parser('prep'); p.add_argument('--urls', required=True); p.add_argument('--links', required=True)
    p.add_argument('--base', required=True); p.add_argument('--body-regex')
    c = sub.add_parser('candidates'); c.add_argument('--base', required=True); c.add_argument('--priority', required=True)
    c.add_argument('--service-regex', required=True); c.add_argument('--gsc-pages'); c.add_argument('--gsc-queries')
    c.add_argument('--gallery-regex'); c.add_argument('--low-value-regex'); c.add_argument('--exclude-regex')
    b = sub.add_parser('build'); b.add_argument('--base', required=True); b.add_argument('--rows', required=True)
    b.add_argument('--client', default=''); b.add_argument('--out', default='internal-linking-audit.xlsx')
    b.add_argument('--findings'); b.add_argument('--note', default='')
    a = ap.parse_args(); {'prep': prep, 'candidates': candidates, 'build': build}[a.cmd](a)
