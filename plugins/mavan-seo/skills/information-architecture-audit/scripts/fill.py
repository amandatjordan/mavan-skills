#!/usr/bin/env python3
"""Fill the IA audit template from work/analysis and run the automated QA checks.

Usage:
  python fill.py --work work --template template.xlsx [--out "Client Information Architecture Audit.xlsx"]
  python fill.py --work work --template template.xlsx --check-only     # QA only, no file

Reads work/config.json, work/decisions.json (recommendations, Read Me notes, tab notes) and the
CSVs analyze.py wrote. Tabs, header text, hover notes, dropdowns, colours and formulas come from
the template; every section is rebuilt by copying the template's own row styles, so sections can
grow past the rows the template reserves without overwriting the next section.
Only values go into data cells. The only formulas in the template are All Pages column C (Area)
and Recommendations column F (Pages); before saving, the script checks that no dropdown, conditional
format or hover note was lost and that no other cell holds a formula.
Exits with status 1 and lists the problems if any check fails; no file is written then.
"""
import argparse
import collections
import os
import re
import sys
from copy import copy

import pandas as pd
from openpyxl import load_workbook
from openpyxl.comments import Comment
from openpyxl.formatting.formatting import ConditionalFormattingList
from openpyxl.worksheet.cell_range import MultiCellRange

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import load_json  # noqa: E402

PRIORITIES = ['High', 'Medium', 'Low', 'Pass', 'Info']
FORMULA_AREAS = ['Duplicate pages', 'Blog URLs', 'Prune candidates', 'Menu cleanup', 'City pages', 'Deep service URLs']
REQUIRED_TOPICS = {
    'junk_pages': 'junk/archive pages', 'blogs_outside_path': 'blogs outside the blog path',
    'services_outside_path': 'services outside the service path', 'duplicates': 'duplicate pages',
    'no_owner': 'searches with no owner page', 'menu_zero_clicks': 'menu pages with 0 clicks',
    'prune': 'prune candidates', 'redirect_links': 'links pointing to redirects/404s',
    'overlapping_browse': 'overlapping ways to browse the same services', 'city_pages': 'city/location pages',
    'service_depth': 'service URL depth', 'traffic_baseline': 'traffic baseline (Info)'}
AREA_FORMULA = ('=IF($A{r}="","",_xlfn.TEXTJOIN(", ",TRUE,IF(OR(AC{r}="Merge then Redirect",AC{r}="Review"),"Duplicate pages",""),'
                'IF(AND(D{r}="Blog",LEFT(B{r},6)<>"/blog/",LEFT(B{r},7)<>"/blogs/",LEFT(B{r},10)<>"/articles/",'
                'LEFT(B{r},6)<>"/news/",LEFT(B{r},11)<>"/resources/"),"Blog URLs",""),IF(AC{r}="Prune candidate",'
                '"Prune candidates",""),IF(AND(O{r}="Yes",N(U{r})=0),"Menu cleanup",""),IF(D{r}="Location / Area",'
                '"City pages",""),IF(AND(OR(D{r}="Service",LEFT(D{r},12)="Service-like"),N(J{r})>=5),'
                '"Deep service URLs",""),AG{r}))')
REC_FORMULA = '=IF(C{r}="","",COUNTIF(\'All Pages\'!C:C,"*"&C{r}&"*"))'
FORMULA_COLUMNS = {('All Pages', 'C'), ('Recommendations', 'F')}  # the only formulas the template holds


def csv(work, name):
    p = os.path.join(work, 'analysis', name)
    if not os.path.exists(p) or os.path.getsize(p) < 2:
        return pd.DataFrame()
    try:
        return pd.read_csv(p, keep_default_na=False)
    except pd.errors.EmptyDataError:
        return pd.DataFrame()


def v(x):
    """Cell value: numbers as numbers, blanks as None."""
    if x is None or (isinstance(x, float) and x != x) or x == '':
        return None
    if isinstance(x, str):
        s = x.strip()
        if re.fullmatch(r'-?\d+', s) and not (len(s) > 1 and s.startswith('0')):
            return int(s)
        if re.fullmatch(r'-?\d+\.\d+', s):
            return float(s)
        return x
    if hasattr(x, 'item'):
        return x.item()
    return x


# ---------------------------------------------------------------- row snapshots
class Proto:
    """A template row's values, styles and hover notes, captured before the sheet is rebuilt."""
    def __init__(self, ws, r, maxc):
        self.cells = []
        for c in range(1, maxc + 1):
            cell = ws.cell(r, c)
            com = cell.comment
            self.cells.append((cell.value, copy(cell._style),
                               (com.text, com.author, com.width, com.height) if com else None))
        self.height = ws.row_dimensions[r].height

    def put(self, ws, r, values=None, keep=False):
        for c, (val, style, com) in enumerate(self.cells, 1):
            cell = ws.cell(r, c)
            cell._style = copy(style)
            if values is not None:
                cell.value = v(values[c - 1]) if c - 1 < len(values) else None
            elif keep:
                cell.value = val
            else:
                cell.value = None
            if com and keep:
                cell.comment = Comment(com[0], com[1], width=com[2], height=com[3])
        if self.height:
            ws.row_dimensions[r].height = self.height


def clear_from(ws, start):
    if ws.max_row >= start:
        ws.delete_rows(start, ws.max_row - start + 1)
    for r in list(ws.row_dimensions):
        if r >= start:
            ws.row_dimensions[r].height = None


def dv_for(ws, col):
    for dv in ws.data_validations.dataValidation:
        if str(dv.sqref).startswith(col):
            return dv
    return None


def set_dv(dv, col, r1, r2):
    if dv is not None:
        dv.sqref = MultiCellRange(f'{col}{r1}:{col}{max(r1, r2)}')


def write_sections(ws, start, sections):
    """sections: dicts with heading/note/header/data Protos, texts, rows and an optional dv column.
    The protos must be captured before this is called."""
    clear_from(ws, start)
    r = start
    for i, s in enumerate(sections):
        if i and s.get('gap', True):
            r += 1
        if s.get('heading'):
            s['heading'].put(ws, r, keep=True)
            if s.get('heading_text'):
                ws.cell(r, 1).value = s['heading_text']
            r += 1
        if s.get('note'):
            s['note'].put(ws, r, keep=True)
            if s.get('note_text') is not None:
                ws.cell(r, 1).value = s['note_text']
            r += 1
        if s.get('header'):
            s['header'].put(ws, r, keep=True)
            r += 1
        first = r
        rows = s.get('rows') or []
        if not rows and s.get('empty_text'):
            s['data'].put(ws, r)
            ws.cell(r, 1).value = s['empty_text']
            r += 1
        for row in rows:
            s['data'].put(ws, r, values=row)
            if s.get('formula'):
                col, f = s['formula']
                ws[f'{col}{r}'] = f.format(r=r)
            r += 1
        if s.get('dv'):
            set_dv(s['dv'][0], s['dv'][1], first, r - 1)
    return r


def sentences(text):
    t = re.sub(r'\b(e\.g|i\.e|etc|vs|approx)\.', '', str(text), flags=re.I)
    t = re.sub(r'\S+\.\S+', 'x', t)  # URLs and file names
    return len(re.findall(r'[.!?](\s|$)', t.strip()))


def how_to_line(ws, sentence=''):
    """Row 1 is the yellow HOW TO USE / WHAT THIS SHOWS line: drop the template's [bracketed]
    instruction and add the site-specific sentence, if any."""
    c = ws['A1']
    base = re.sub(r'\s*\[[^\]]*\]', '', str(c.value or '')).rstrip()
    if sentence and sentence.strip() not in base:
        base = f'{base} {sentence.strip()}'
    c.value = base


# ---------------------------------------------------------------- integrity
def snapshot(wb):
    """Dropdowns, conditional format rules and hover notes in the workbook, as multisets, plus any cell
    that holds a formula outside the two formula columns."""
    dv, cf, notes, stray = collections.Counter(), collections.Counter(), collections.Counter(), []
    for ws in wb.worksheets:
        for d in ws.data_validations.dataValidation:
            dv[(ws.title, d.type, str(d.formula1))] += 1
        for rng in ws.conditional_formatting:
            for rule in rng.rules:
                cf[(ws.title, rule.type, rule.operator, str(rule.formula))] += 1
        for row in ws.iter_rows():
            for c in row:
                if c.comment:
                    notes[(ws.title, c.comment.text)] += 1
                if isinstance(c.value, str) and c.value.startswith('=') and (ws.title, c.column_letter) not in FORMULA_COLUMNS:
                    stray.append(f'{ws.title}!{c.coordinate}')
    return {'dropdowns': dv, 'conditional formats': cf, 'hover notes': notes}, stray


def integrity_errors(before, after, stray_after):
    errs = []
    for what in before:
        lost = before[what] - after[what]
        extra = after[what] - before[what]
        if lost:
            errs.append(f'{sum(lost.values())} {what} from the template are missing after filling, e.g. '
                        f'{next(iter(lost))[0]}: {str(next(iter(lost))[1:])[:80]}')
        if extra:
            errs.append(f'{sum(extra.values())} {what} appeared that the template did not have, e.g. '
                        f'{next(iter(extra))[0]}: {str(next(iter(extra))[1:])[:80]}')
    if stray_after:
        errs.append('Formulas in cells that are not Area (All Pages column C) or Pages (Recommendations column F): '
                    + ', '.join(stray_after[:8]))
    return errs


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--work', default='work')
    ap.add_argument('--template', required=True)
    ap.add_argument('--out')
    ap.add_argument('--check-only', action='store_true')
    a = ap.parse_args()
    W = a.work
    cfg = load_json(os.path.join(W, 'config.json'), {})
    dec = load_json(os.path.join(W, 'decisions.json'), {})
    sm = load_json(os.path.join(W, 'analysis', 'summary.json'), {})
    if not sm:
        sys.exit('Run analyze.py first: work/analysis/summary.json is missing.')
    errors, warns = [], []
    P = csv(W, 'all_pages.csv')
    RD = csv(W, 'redirects.csv')

    # ---------------- recommendations checks
    recs = dec.get('recommendations', [])
    extra_areas = set()
    for x in P.extra_area.tolist() if len(P) else []:
        extra_areas |= {y.strip() for y in str(x).split(',') if y.strip()}
    area_count = {}
    if len(P):
        num = lambda s: pd.to_numeric(s, errors='coerce').fillna(0)
        area_count = {
            'Duplicate pages': int(P.action.isin(['Merge then Redirect', 'Review']).sum()),
            'Blog URLs': int(((P.type == 'Blog') & ~P.path.str.startswith(('/blog/', '/blogs/', '/articles/', '/news/', '/resources/'))).sum()),
            'Prune candidates': int((P.action == 'Prune candidate').sum()),
            'Menu cleanup': int(((P.in_nav == 'Yes') & (num(P.clicks) == 0)).sum()),
            'City pages': int((P.type == 'Location / Area').sum()),
            'Deep service URLs': int((((P.type == 'Service') | P.type.str.startswith('Service-like')) & (num(P.path_depth) >= 5)).sum()),
        }
        for ar in extra_areas:
            area_count[ar] = int(P.extra_area.str.contains(re.escape(ar)).sum())
    # The Pages formula is COUNTIF(Area, "*name*"): case-insensitive substring. A name inside another
    # (e.g. "Service URLs" inside "Deep service URLs") would count both.
    used = [rc.get('area', '') for rc in recs if rc.get('area') and 'pages' not in rc]
    all_names = set(area_count)
    for ar in used:
        clash = [o for o in all_names if o != ar and ar.lower() in o.lower() and area_count.get(o)]
        if clash:
            errors.append(f'Area "{ar}" is part of {clash}, so its Pages count would include those rows too. Rename it.')
    if not recs:
        errors.append('decisions.json has no recommendations.')
    covered = set()
    for i, rc in enumerate(recs, 1):
        tag = f'Recommendation {i} ({rc.get("area", "?")})'
        if rc.get('priority') not in PRIORITIES:
            errors.append(f'{tag}: priority must be one of {PRIORITIES}, got {rc.get("priority")!r}.')
        for f in ('area', 'problem', 'do', 'where'):
            if not str(rc.get(f, '')).strip():
                errors.append(f'{tag}: "{f}" is empty.')
        ar = rc.get('area', '')
        if ar and ar not in area_count and 'pages' not in rc:
            errors.append(f'{tag}: area "{ar}" is not tagged on All Pages (formula areas: {FORMULA_AREAS}; '
                          f'extra areas: {sorted(extra_areas)}). Add "pages": <number> or tag the rows via extra_areas.')
        prob = str(rc.get('problem', ''))
        if sentences(prob) > 1:
            warns.append(f'{tag}: Problem should be one sentence.')
        if sentences(rc.get('do', '')) > 2:
            warns.append(f'{tag}: Do this should be one or two short sentences.')
        txt = (prob + ' ' + str(rc.get('do', ''))).lower()
        if re.search(r'(add|more|build)\s+(internal\s+)?links?\s+(to|into|pointing (to|at))\s+(the\s+|these\s+|this\s+|your\s+)?(blog|post)', txt):
            errors.append(f'{tag}: recommends adding links to blog posts, which this audit must not do.')
        covered |= set(rc.get('covers', []))
    missing = [REQUIRED_TOPICS[t] for t in REQUIRED_TOPICS if t not in covered]
    if recs and missing:
        errors.append('Recommendations do not cover (tag each row with "covers"): ' + '; '.join(missing))

    # ---------------- redirect checks
    if len(RD):
        planned = RD[RD.kind == 'Planned']
        for r in planned.itertuples():
            if str(r.target_live) != 'Yes':
                errors.append(f'Redirect {r.source} -> {r.target}: target is not a live 200 page.')
            if str(r.target_redirected) == 'Yes':
                errors.append(f'Redirect {r.source} -> {r.target}: target is itself being redirected.')
            if str(r.self_match) == 'Yes':
                errors.append(f'Wildcard {r.source}: its target {r.target} matches its own pattern.')
        for r in RD[RD.kind == 'Existing redirect still linked'].itertuples():
            if str(r.target_live) != 'Yes':
                warns.append(f'Existing redirect {r.source} ends at {r.target}, which is not a live page.')
    if sm.get('nav_share', 0) and sm['nav_share'] >= 0.95:
        errors.append('95%+ of links are Nav: link classification failed. Fix it in prep.py before filling.')
    if not sm.get('mapping_confirmed'):
        warns.append('The folder-to-type mapping was not confirmed by a person; page types are unverified.')
    if sm.get('unmapped_folders'):
        warns.append('Folders missing from the confirmed mapping: ' + ', '.join(sm['unmapped_folders'][:10]))
    if sm.get('pairs_to_check'):
        warns.append(f'{sm["pairs_to_check"]} duplicate pair(s) still marked "check"; they are treated as not duplicates.')
    if sm.get('has_gsc') and not sm.get('pair_queries_loaded') and sm.get('pair_queries_needed'):
        warns.append('Top shared queries were not loaded (ia_pair_queries); that column will be blank.')

    # ---------------- write
    wb = load_workbook(a.template)
    before, stray_before = snapshot(wb)
    if stray_before:
        errors.append('The template has formulas outside All Pages column C and Recommendations column F '
                      f'({", ".join(stray_before[:6])}); use the current template.')
    client = cfg.get('client', 'Client')

    # Read Me
    ws = wb['Read Me']
    ws['A1'] = str(ws['A1'].value).replace('{Company}', client)
    meta = sm.get('gsc_meta', {})
    rd_notes = dec.get('readme', {})
    uc = sm.get('url_counts') or {}
    lc = sm.get('link_classes') or {}
    if sm.get('has_gsc'):
        win = f"{meta.get('window_start')} to {meta.get('window_end')}"
        window_text = f'{win}: the last 6 full calendar months.'
        totals = (f"{int(float(meta.get('clicks', 0))):,} clicks, {int(float(meta.get('impressions', 0))):,} impressions, "
                  f"{meta.get('page_rows')} page rows, {meta.get('page_query_rows')} page+query rows (after joining URL variants). "
                  f"Brand searches excluded from query analysis: {int(float(meta.get('brand_query_impr', 0))):,} impressions.")
        ga4 = (f"{meta.get('ga4_property')}, dimension landingPage, metrics sessions, keyEvents "
               f"({meta.get('ga4_rows')} landing pages)" if meta.get('ga4_property') else 'GA4 not available: ' +
               rd_notes.get('ga4_missing', 'no GA4 property was found for this domain.'))
        prop = meta.get('site_url', '')
    else:
        window_text = totals = prop = 'Search Console data was not available for this run.'
        ga4 = 'Not available.'
    utm = sm.get('utm_home') or []
    readme = {
        4: f"{client} - {cfg.get('domain')} - {cfg.get('city', '')}",
        5: rd_notes.get('other_cities') or (', '.join(cfg.get('other_cities', [])) or 'none'),
        6: ', '.join(cfg.get('brand_terms', [])) or 'none',
        7: f"{sm.get('crawler')}, export date {cfg.get('crawl_date', 'not given')}, {sm.get('urls_total')} URLs: " +
           ', '.join(f'{k} {v}' for k, v in sorted(uc.items())) + f"; {sm.get('live_pages')} live pages on All Pages.",
        8: prop,
        9: window_text,
        10: totals,
        11: ('; '.join(utm) + '. Its clicks are counted in the homepage row because tracking parameters are removed when URLs are joined.')
            if utm else 'none found',
        12: ga4,
        13: 'Every URL lowercased, query strings (?...) and fragments (#...) removed, forced to https and non-www, '
            'trailing slash added (not on files). Rows that became the same URL were added together.',
        14: (('Folder-to-type mapping confirmed by a person: ' if sm.get('mapping_confirmed') else
              'Folder-to-type mapping NOT confirmed (folder types were guessed from common folder names): ') +
             ('; '.join(f'/{f.strip("/")}/ = {t}' for f, t in (sm.get('folder_mapping') or {}).items()) or 'none') + '. '
             'In order: 1 "/" = Home. 2 ?page= or /page/N/ = Pagination; other query strings = Query-string URL. '
             '3 the index page of a mapped folder = that folder\'s hub type; any other page in a mapped folder takes the '
             'folder\'s type. 4 folders mapped as Mixed and root-level pages use content rules: utility wording (privacy, '
             'terms, HIPAA, thank-you, sitemap, test) = Utility / Policy; company words = Company / Info; business words '
             '+ city only = Location / Area; question-word slug, 6+ word slug, blog wording (vs, benefits, guide, tips, '
             'cost, explained, what to expect, reasons, aftercare, before and after, myths) or a question H1 = Blog; '
             'other short slugs = Service-like (no folder); anything left = Unclassified. Reviewer corrections are marked '
             '"Set by reviewer".'),
        15: 'Classified from the CSS path / link path, not the crawler location label: menu, nav, header, mega menu and '
            'accordion containers = Nav; footer = Footer; breadcrumbs = Breadcrumb; everything else = Body. Fragments '
            'removed; each link counted once per linking page + target page. Counts: ' +
            ', '.join(f'{k} {v:,}' for k, v in lc.items()) + f" (from {sm.get('link_rows', 0):,} link rows).",
        16: 'Pairs of content pages (service, service-like, concern, blog, location) sharing non-brand searches: per '
            'shared search the smaller of the two pages\' impressions, summed, divided by the smaller page\'s total '
            'non-brand query impressions. Kept at 0.30+ with 2+ shared searches; also pages with the same H1 after '
            'removing cities and filler words. A true duplicate needs the same treatment. High = 0.50+ with 3+ shared '
            'searches, Medium = other search overlap, Low = same H1 only or a pair that could not be judged. Keeper: '
            'service page, then concern, then blog; main city first; then most impressions. Searches with no owner '
            'page: 3+ treatments, no page holding half the impressions, and at least 100 impressions in the window.',
        17: '; '.join([x for x in [*sm.get('notes', []), rd_notes.get('missing', '')] if x]) or 'Nothing missing.',
        18: rd_notes.get('side_observations', 'None.'),
    }
    for r, txt in readme.items():
        ws.cell(r, 2).value = txt
    # Recommendations
    ws = wb['Recommendations']
    order = {p: i for i, p in enumerate(PRIORITIES)}
    recs_sorted = sorted(recs, key=lambda x: order.get(x.get('priority'), 9))
    pr = Proto(ws, 5, 9)
    rows = [[i, x.get('priority'), x.get('area'), x.get('problem'), x.get('do'), None, x.get('where'), None, None]
            for i, x in enumerate(recs_sorted, 1)]
    dvh = dv_for(ws, 'H')
    end = write_sections(ws, 5, [{'data': pr, 'rows': rows, 'formula': ('F', REC_FORMULA), 'dv': (dvh, 'H')}])
    for i, x in enumerate(recs_sorted):
        if 'pages' in x:
            ws.cell(5 + i, 6).value = x['pages']
    cf_old = ws.conditional_formatting
    ws.conditional_formatting = ConditionalFormattingList()
    for rng in cf_old:
        for rule in rng.rules:
            ws.conditional_formatting.add(f'B5:B{max(40, end)}', rule)

    # All Pages
    ws = wb['All Pages']
    pa = Proto(ws, 3, 33)
    has_past = len(P) and (P.past_decision != '').any()
    rows = []
    for r in P.itertuples():
        rows.append([r.url, r.path, None, r.type, r.why, r.topic, r.status, r.indexable, r.crawl_depth, r.path_depth,
                     r.title, r.h1, r.h2, r.words, r.in_nav, r.nav_in, r.menu_links, r.body_in, r.body_in_arch,
                     r.footer_in, r.clicks, r.impressions, r.position, r.top_query, r.top_query_impr, r.ga4_sessions,
                     r.ga4_key_events, r.dup_note, r.action, r.target, r.final_action, r.final_target,
                     r.extra_area])
    write_sections(ws, 3, [{'data': pa, 'rows': rows, 'formula': ('C', AREA_FORMULA)}])
    if has_past:
        hdr = Proto(ws, 2, 33)
        for j, name in enumerate(['Past decision', 'Past target', 'Match?']):
            c = ws.cell(2, 34 + j); c._style = copy(hdr.cells[28][1]); c.value = name
        for i, r in enumerate(P.itertuples(), 3):
            for j, val in enumerate([r.past_decision, r.past_target, r.match]):
                c = ws.cell(i, 34 + j); c._style = copy(pa.cells[28][1]); c.value = v(val)

    # Page Type Summary
    ws = wb['Page Type Summary']
    S = csv(W, 'page_type_summary.csv')
    ps, t1, t2, t3 = Proto(ws, 3, 13), Proto(ws, 19, 13), Proto(ws, 20, 13), Proto(ws, 21, 13)
    rows = [[r.type, r.pages, r.pct_indexable, r.clicks, r.share, r.impressions, r.pct_zero, r.med_body, r.med_nav,
             r.med_depth, r.med_words, r.ga4_sessions, r.ga4_key_events] for r in S.itertuples()]
    src = (f"Source: {sm.get('crawler')} crawl; Search Console {meta.get('window_start', '')} to {meta.get('window_end', '')}"
           f"{'; GA4 ' + meta.get('ga4_property') if meta.get('ga4_property') else ''}. Body inlinks exclude links from archive pages.")
    write_sections(ws, 3, [
        {'data': ps, 'rows': rows},
        {'data': t1, 'rows': [[t1.cells[0][0], int(float(meta.get('clicks', 0))) if meta else None]]},
        {'data': t2, 'rows': [[t2.cells[0][0], int(float(meta.get('impressions', 0))) if meta else None]], 'gap': False},
        {'data': t3, 'rows': [[src]], 'gap': False}])

    # Duplicate Intent
    ws = wb['Duplicate Intent']
    protos = {r: Proto(ws, r, 17) for r in (2, 3, 4, 5, 46, 47, 48, 79, 80, 81, 83, 84, 85, 86, 100, 101, 102, 103)}
    G = csv(W, 'dup_groups.csv')
    rows = [[r.group, r.treatment, r.group_confidence, r.keeper, r.page, r.page_type, r.role, r.clicks, r.impressions,
             r.best_overlap, r.shared_queries, r.shared_impr, r.evidence, r.top_shared, r.outranks, r.action, None]
            for r in G.itertuples()]
    E = csv(W, 'dup_pairs.csv')
    erows = [[r.treatment, r.page_a, r.page_b, r.confidence, r.overlap, r.shared_queries, r.shared_impr, r.evidence,
              r.top_shared] for r in E.itertuples()]
    C = csv(W, 'city_variants.csv')
    crows = [[r.page_a, r.page_b, r.overlap, r.shared_queries, r.shared_impr, r.top_shared, r.note] for r in C.itertuples()]
    N = csv(W, 'no_owner.csv')
    nrows = [[r.query, r.impressions, r.pages, r.treatments, r.top_pages, r.suggested] for r in N.itertuples()]
    H = csv(W, 'home_overlap.csv')
    hrows = [[r.page, r.page_type, r.share, r.shared_queries, r.nb_impr] for r in H.itertuples()]
    home_note = ((f"Homepage data includes the GBP/UTM link ({utm[0]}). " if utm else 'No GBP/UTM homepage URL found. ') +
                 'Top non-brand homepage queries: ' + (meta.get('home_top_queries') or 'none'))
    dq = dv_for(ws, 'Q')
    write_sections(ws, 2, [
        {'heading': protos[2], 'note': protos[3], 'header': protos[4], 'data': protos[5], 'rows': rows,
         'dv': (dq, 'Q'), 'empty_text': 'No duplicate groups found for this site.'},
        {'heading': protos[46], 'header': protos[47], 'data': protos[48], 'rows': erows, 'empty_text': 'None.'},
        {'heading': protos[79], 'header': protos[80], 'data': protos[81], 'rows': crows,
         'empty_text': 'No city variants found.'},
        {'heading': protos[83], 'note': protos[84], 'header': protos[85], 'data': protos[86], 'rows': nrows,
         'empty_text': 'None found.'},
        {'heading': protos[100], 'note': protos[101], 'note_text': home_note, 'header': protos[102],
         'data': protos[103], 'rows': hrows, 'empty_text': 'None found.'}])

    # Path Consistency
    ws = wb['Path Consistency']
    pp = {r: Proto(ws, r, 7) for r in (2, 3, 4, 26, 27, 28, 29)}
    PP = csv(W, 'path_patterns.csv')
    FL = csv(W, 'path_flags.csv')
    flag_counts = FL.flag.str.replace(r' \(.*\)$', '', regex=True).value_counts().to_dict() if len(FL) else {}
    path_summary = dec.get('path_summary') or ('; '.join(f'{n} x {k}' for k, n in flag_counts.items()) or
                                                'No path problems found.')
    dg = dv_for(ws, 'G')
    write_sections(ws, 2, [
        {'heading': pp[2], 'header': pp[3], 'data': pp[4], 'dv': (dg, 'G'),
         'rows': [[r.type, r.pattern, r.pages, r.pct, r.examples, r.main, None] for r in PP.itertuples()]},
        {'heading': pp[26], 'note': pp[27], 'note_text': path_summary, 'header': pp[28], 'data': pp[29],
         'rows': [[r.url, r.type, r.flag] for r in FL.itertuples()], 'empty_text': 'No flags.'}])

    # Link Flow
    ws = wb['Link Flow']
    pl, pr18 = Proto(ws, 4, 18), Proto(ws, 18, 18)
    LF = csv(W, 'link_flow.csv')
    cols = ['Home', 'Service hub', 'Service', 'Service-like', 'Concern hub', 'Concern', 'Location / Area', 'Blog hub',
            'Blog', 'Company / Info', 'Utility / Policy', 'Pagination', 'Query-string URL', 'Redirect (3xx)']
    rows = [[r['type']] + [r[c] for c in cols] + [r['total'], r['pct_commercial'], r['pct_redirect']]
            for _, r in LF.iterrows()]
    other = int(LF.other.sum()) if len(LF) else 0
    tot = int(LF.total.sum()) if len(LF) else 0
    flow_note = dec.get('link_flow_note') or (
        f"{tot:,} body links in total. Redirect (3xx) also counts links to 404 pages. " +
        (f"{other:,} links go to page types with no column here (archives, gallery items, other folders); they are in "
         "the totals." if other else ''))
    write_sections(ws, 4, [{'data': pl, 'rows': rows, 'empty_text': 'No body links found.'},
                           {'data': pr18, 'rows': [[f'Reading it: {flow_note}']]}])

    # Retarget Links
    ws = wb['Retarget Links']
    pr_ = {r: Proto(ws, r, 9) for r in (2, 3, 5, 6, 7)}
    RT, RR = csv(W, 'retarget.csv'), csv(W, 'retarget_ref.csv')
    ref_rows = [[r.src, r.src_type, r.blog, r.anchor, r.image, r.named, r.info, r.blog_clicks, r.match] for r in RR.itertuples()]
    if len(RT):
        res_rows = [[r.src, '', r.blog, r.anchor, '', '', '', '', r.new_target] for r in RT.itertuples()]
        typ = dict(zip(P.path, P.type)) if len(P) else {}
        for row in res_rows:
            row[1] = typ.get(row[0], '')
        sec1 = {'heading': pr_[2], 'header': pr_[6], 'data': pr_[7], 'rows': res_rows}
    else:
        sec1 = {'heading': pr_[2], 'note': pr_[3], 'note_text': sm.get('retarget_none_reason') or 'None found.',
                'data': pr_[7], 'rows': []}
    write_sections(ws, 2, [sec1, {'heading': pr_[5], 'header': pr_[6], 'data': pr_[7], 'rows': ref_rows,
                                  'empty_text': 'No body links point at blogs with 2 or fewer clicks.'}])

    # Blog to Service Links
    ws = wb['Blog to Service Links']
    pb = Proto(ws, 3, 17)
    B = csv(W, 'blog_to_service.csv')
    rows = [[r.blog, r.title, r.h1, r.topic, r.service, r.status, r.old_url, r.clicks, r.impressions, r.body_in,
             r.arch_in, r.feed_in, r.other_in, r.body_out, r.action, r.risk, None] for r in B.itertuples()]
    write_sections(ws, 3, [{'data': pb, 'rows': rows, 'dv': (dv_for(ws, 'Q'), 'Q'), 'empty_text': 'No blogs found.'}])

    # Nav Pages
    ws = wb['Nav Pages']
    pn = Proto(ws, 3, 17)
    NV = csv(W, 'nav_pages.csv')
    rows = [[r.page, r.type, r.title, r.h1, r.crawl_depth, r.nav_in, r.menu_links, r.body_in, r.footer_in, r.words,
             r.clicks, r.impressions, r.position, r.top_query, r.indexable, r.suggested, None]
            for r in NV.itertuples()] if len(NV) and 'type' in NV else []
    write_sections(ws, 3, [{'data': pn, 'rows': rows, 'dv': (dv_for(ws, 'Q'), 'Q'), 'empty_text': 'No menu links found.'}])

    # Redirects
    ws = wb['Redirects']
    prd = Proto(ws, 3, 14)
    rows = [[r.kind, r.source, r.target, r.when, r.detail, r.status, r.body_in, r.in_nav, r.footer_in, r.pages_linking,
             r.examples, r.target_live, r.target_redirected, None] for r in RD.itertuples()]
    write_sections(ws, 3, [{'data': prd, 'rows': rows, 'dv': (dv_for(ws, 'N'), 'N'), 'empty_text': 'No redirects needed.'}])

    notes = dec.get('tab_notes', {})
    for wsx in wb.worksheets:
        if wsx.title != 'Read Me':
            how_to_line(wsx, notes.get(wsx.title, ''))

    # ---------------- placeholder scan
    for wsx in wb.worksheets:
        for row in wsx.iter_rows():
            for c in row:
                if isinstance(c.value, str) and not c.value.startswith('=') and re.search(r'\[[A-Z][^\]]{6,}\]', c.value):
                    errors.append(f'{wsx.title}!{c.coordinate} still has a template placeholder: {c.value[:70]}')
                if isinstance(c.value, str) and '{Company}' in c.value:
                    errors.append(f'{wsx.title}!{c.coordinate} still says {{Company}}')

    # ---------------- integrity: nothing the template carried may be lost
    after, stray_after = snapshot(wb)
    errors.extend(integrity_errors(before, after, stray_after))

    # ---------------- report
    print('Area counts on All Pages (what the Recommendations Pages column will show):')
    for ar, n in sorted(area_count.items(), key=lambda x: -x[1]):
        print(f'  {ar}: {n}')
    for rc in recs_sorted:
        if 'pages' in rc:
            print(f'  {rc.get("area")}: {rc["pages"]} (typed in, not tagged on All Pages)')
    for w in warns:
        print('WARNING:', w)
    if errors:
        print('\nPROBLEMS (fix these, then run again):')
        for e in errors:
            print(' -', e)
        sys.exit(1)
    if a.check_only:
        print('\nAll checks passed.')
        return
    out = a.out or f'{client} Information Architecture Audit.xlsx'
    wb.save(out)
    print(f'\nAll checks passed (template dropdowns, conditional formats and hover notes intact). Saved {out}')


if __name__ == '__main__':
    main()
