#!/usr/bin/env python3
"""Service Gap Audit.

  python build.py inspect --crawl export.csv --cities "City, City" --practice medspa|surgical|both --out work
  python build.py build   --crawl export.csv --pages work/pages.tsv --cities "City, City" --practice ... \
                          --out DOMAIN_service_gap_audit.xlsx [--gsc GSC_SHEET [--gsc-tab TAB]] [--taxonomy FILE]

inspect: maps the export's columns, suggests a page type and taxonomy node for every live page from its URL,
         title and H1, and writes work/pages.tsv. The AI corrects the `type` and `node` columns there: page type
         is never decided from a fixed folder name, because every site is built differently.
build:   classifies the competitor keywords with the taxonomy, works out coverage, offered status and gaps for
         every taxonomy node, assigns keywords to recommended pages, adds Search Console impressions when --gsc is
         given, builds the workbook and runs the QA checks (PASS / FAIL).
Every number, count and join happens here. Rules: references/rules.md. Workbook: references/workbook.md.
"""
import argparse
import csv
import difflib
import os
import re
import sys
from collections import defaultdict
from urllib.parse import unquote, urlsplit

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

HERE = os.path.dirname(os.path.abspath(__file__))
DEFAULT_TAXONOMY = os.path.join(HERE, "..", "references", "service-taxonomy.xlsx")

THIN, LIGHT = 300, 600                      # content words
HIGH, MEDIUM = 5000, 2000                   # volume tiers
MODIFIER_SECTION_MIN = 1000                 # page-brief modifier sections
PAGE_TYPES = ["Service", "Concern", "Gallery", "Blog", "Location", "Utility"]
NEW_PREFIX = "NEW:"                         # node value for a brand/device missing from the taxonomy: "NEW: Name > Service"
NO_KW_TEXT = "None in competitor keyword list (check GSC / keyword tool)"

COLS = {  # canonical name: accepted export headers (lowercase)
    "url": ["url", "address"],
    "status": ["http status code", "status code"],
    "indexable": ["indexable status", "indexability status", "indexability"],
    "title": ["title", "title 1", "page title"],
    "meta": ["meta description", "meta description 1"],
    "h1": ["h1", "h1-1"],
    "h2": ["h2", "h2-1"],
    "h2b": ["h2 (second)", "h2-2"],
    "words": ["no. content words", "word count", "no. words"],
    "canonical": ["canonical url", "canonical link element 1", "canonical"],
}
REQUIRED = ["url", "status", "title", "h1", "words"]

LABELS = [  # Why Recommended, in sort order
    "Existing page, not indexable",
    "Existing page, thin content",
    "Offered (before & after photos), no page",
    "Offered (named on a concern page), no page",
    "Likely offered (mentioned on site), no page",
    "Only has a brand/product page",
    "Only has Sub-Service pages",
    "Only covered in blog posts",
    "Existing page, light content",
    "Confirm with client (parent Service offered)",
    "Confirm with client (no sign it is offered)",
]

ARIAL = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3A5F")
YELLOW = PatternFill("solid", fgColor="FFF2CC")
PRIO_FILL = {"P1": "F8CBAD", "P2": "FFE699", "P3": "E2EFDA", "Ask": "DDEBF7"}


# ------------------------------------------------------------------ helpers
def read_table(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        return pd.read_excel(path, dtype=str).fillna("")
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", encoding_errors="replace")


def page_path(u):
    p = unquote(urlsplit(u.strip()).path or "/")
    return p if p else "/"


def path_key(u):
    return page_path(u).lower().rstrip("/") or "/"


def to_int(x):
    try:
        return int(float(str(x).replace(",", "").strip() or 0))
    except ValueError:
        return 0


def slugtext(p):
    return re.sub(r"[-_/.]+", " ", p.lower())


def clean(s):
    s = str(s or "").strip()
    return "" if s in ("--", "nan", "None") else s


def cities_list(s):
    return [c.strip().lower() for c in (s or "").split(",") if c.strip()]


# ------------------------------------------------------------------ taxonomy
class Taxonomy:
    def __init__(self, path):
        x = pd.read_excel(path, sheet_name=None, dtype=str)
        for need in ("Taxonomy", "Filters", "All Keywords"):
            if need not in x:
                sys.exit(f"{path}: no '{need}' tab. Use the bundled taxonomy (references/service-taxonomy.xlsx).")
        self.version = str(list(x["Read Me"].columns)[0]) if "Read Me" in x else os.path.basename(path)
        t = x["Taxonomy"].fillna("")
        self.nodes = []
        for i, r in t.iterrows():
            lvl = r["Level"].strip()
            name = (r["Brand / Device"] or r["Sub-Service"] or r["Concern"] or r["Service"]).strip() if lvl != "Concern" else r["Concern"].strip()
            if lvl == "Sub-Service":
                name = r["Sub-Service"].strip()
            elif lvl == "Service":
                name = r["Service"].strip()
            rx = r["Match Terms (regex)"].strip()
            if not rx:
                continue
            subs = [s.strip() for s in r["Sub-Service"].split(";") if s.strip()] if lvl == "Brand / Device" else []
            implies = [s.strip() for s in r["Implies Offered"].split(";") if s.strip()]
            self.nodes.append(dict(idx=i, level=lvl, name=name, service=r["Service"].strip(), category=r["Category"].strip(),
                                   subs=subs, implies=implies, practice=r["Practice Type"].strip(),
                                   note=r["Status / Notes"].strip(), rx=re.compile(rx, re.I)))
        self.by_level = defaultdict(list)
        for n in self.nodes:
            self.by_level[n["level"]].append(n)
        self.key = {(n["level"], n["name"], n["service"]): n for n in self.nodes}
        self.names = {}
        for n in self.nodes:
            if n["level"] != "Concern":
                self.names.setdefault(n["name"].lower(), n)
        f = x["Filters"].fillna("")
        self.filters = {r["Filter"].strip(): r["Pattern (regex) or word list"] for _, r in f.iterrows()}
        self.fx = {}
        for k in ("Celebrity / person", "Competitor brand", "Off-topic", "Practice-type / generic", "Near me"):
            if k in self.filters:
                self.fx[k] = re.compile(self.filters[k], re.I)
        self.mods = {r["Filter"]: re.compile(r["Pattern (regex) or word list"], re.I)
                     for _, r in f.iterrows() if r["Type"].strip() == "Modifier"}
        self.kw = x["All Keywords"].fillna("")
        self.kw["Volume"] = self.kw["Volume"].map(to_int)

    def best(self, text, level, skip=lambda n: False):
        """Longest match at this level; ties go to the earlier taxonomy row."""
        hit = None
        for n in self.by_level[level]:
            if skip(n):
                continue
            m = n["rx"].search(text)
            if m:
                score = (m.end() - m.start(), -n["idx"])
                if hit is None or score > hit[0]:
                    hit = (score, n)
        return hit[1] if hit else None

    def classify_keyword(self, k):
        k = k.lower()
        for f in ("Celebrity / person", "Competitor brand", "Off-topic"):
            if f in self.fx and self.fx[f].search(k):
                return ("Exclude", "", "", "", "", "")
        b, s, sv = self.best(k, "Brand / Device"), self.best(k, "Sub-Service"), self.best(k, "Service")
        service = sub = brand = ""
        if b:
            service, brand = b["service"], b["name"]
        if s and (not service or s["service"] == service):
            service, sub = s["service"], s["name"]
        if b and not sub and len(b["subs"]) == 1:
            sub = b["subs"][0]
        if not service and sv:
            service = sv["service"]
        if service:
            return ("Service", "Brand / Device" if brand else "Sub-Service" if sub else "Service", service, sub, brand, "")
        c = self.best(k, "Concern")
        if c:
            return ("Concern", "Concern", "", "", "", c["name"])
        if "Practice-type / generic" in self.fx and self.fx["Practice-type / generic"].search(k):
            return ("Generic", "", "", "", "", "")
        return ("Unmapped", "", "", "", "", "")

    def page_node(self, text):
        """Most specific taxonomy node named in a page's URL + title + H1 (Concerns ignored)."""
        for lvl in ("Brand / Device", "Sub-Service", "Service"):
            n = self.best(text, lvl)
            if n:
                return n
        return None

    def lookup(self, name):
        return self.names.get(name.strip().lower())


def fits_practice(node_practice, practice):
    if practice == "both" or node_practice in ("Both", ""):
        return True
    return (practice == "medspa" and node_practice == "Medspa") or (practice == "surgical" and node_practice == "Surgical")


def flagged_do_not(note):
    n = note.lower()
    return "do not flag" in n or "not us-approved" in n


# ------------------------------------------------------------------ crawl
def load_crawl(path, overrides):
    df = read_table(path)
    lookup = {c.strip().lower(): c for c in df.columns}
    cols = {}
    for key, names in COLS.items():
        o = overrides.get(key)
        if o:
            if o.strip().lower() not in lookup:
                sys.exit(f"--{key}-col '{o}' is not a column. Columns: {list(df.columns)}")
            cols[key] = lookup[o.strip().lower()]
            continue
        cols[key] = next((lookup[n] for n in names if n in lookup), None)
    missing = [k for k in REQUIRED if not cols[k]]
    if missing:
        sys.exit(f"Could not find column(s) {missing}. Columns: {list(df.columns)}. Re-run with --<name>-col.")
    rows = []
    for _, r in df.iterrows():
        url = clean(r[cols["url"]])
        if not url:
            continue
        rows.append(dict(url=url, path=page_path(url), key=path_key(url), status=to_int(r[cols["status"]]),
                         indexable=clean(r[cols["indexable"]]) if cols["indexable"] else "",
                         title=clean(r[cols["title"]]), h1=clean(r[cols["h1"]]),
                         h2=clean(r[cols["h2"]]) if cols["h2"] else "", h2b=clean(r[cols["h2b"]]) if cols["h2b"] else "",
                         meta=clean(r[cols["meta"]]) if cols["meta"] else "", words=to_int(r[cols["words"]]),
                         canonical=clean(r[cols["canonical"]]) if cols["canonical"] else ""))
    return rows, cols


def indexability(rows):
    """Adds 'index' to every row: Indexable, Canonical loop, or the crawler's status."""
    by_url = {r["url"].strip(): r for r in rows}
    by_key = defaultdict(list)
    for r in rows:
        by_key[r["key"]].append(r)
    for r in rows:
        st = r["indexable"]
        if r["status"] != 200:
            r["index"] = f"Not 200 ({r['status']})"
            continue
        if st.lower().startswith("indexable") or st == "":
            r["index"] = "Indexable"
            continue
        loop = False
        if "canonical" in st.lower():
            tgt = by_url.get(r["canonical"]) if r["canonical"] else None
            if tgt is not None:
                loop = 300 <= tgt["status"] < 400
            else:  # no canonical column: same path with the other trailing-slash form redirects
                loop = any(o is not r and 300 <= o["status"] < 400 for o in by_key[r["key"]])
        r["index"] = "Canonical loop (canonical target redirects)" if loop else (st or "Not indexable")


# ------------------------------------------------------------------ inspect
GALLERY_HINT = re.compile(r"before[-_ ]?(and[-_ ]?|&[-_ ]?)?after|gallery|galleries|\bphotos?\b|\bresults\b", re.I)
BLOG_HINT = re.compile(r"\b(blog|news|articles?|posts?|insights|resources)\b|/\d{4}/\d{2}/|/page/\d+", re.I)
UTILITY_HINT = re.compile(r"\b(contact|about|team|staff|privacy|terms|policy|careers|sitemap|thank|login|cart|checkout|"
                          r"account|faq|reviews?|testimonials|specials?|promos?|financing|membership|refill|forms?|book|"
                          r"appointment|search|404)\b", re.I)


def suggest(r, tax, cities):
    text = " | ".join([slugtext(r["path"]), r["title"].split("|")[0], r["h1"]])
    node = tax.page_node(text)
    concern = tax.best(" | ".join([r["title"].split("|")[0], r["h1"]]), "Concern")
    city = next((c for c in cities if c in (slugtext(r["path"]) + " " + r["h1"]).lower()), "")
    if GALLERY_HINT.search(slugtext(r["path"])) and r["words"] < THIN:
        t = "Gallery"
    elif BLOG_HINT.search(r["path"]):
        t = "Blog"
    elif node:
        t = "Service"
    elif concern:
        t = "Concern"
    elif UTILITY_HINT.search(slugtext(r["path"]) + " " + r["title"].lower()) or r["path"] == "/":
        t = "Utility"
    elif city:
        t = "Location"
    else:
        t = "Utility"
    return t, (node["name"] if node else ""), (node["service"] if node else ""), city


def cmd_inspect(a):
    tax = Taxonomy(a.taxonomy)
    rows, cols = load_crawl(a.crawl, overrides_from(a))
    indexability(rows)
    cities = cities_list(a.cities)
    print("Columns used: " + ", ".join(f"{k} = '{v}'" for k, v in cols.items() if v))
    absent = [k for k, v in cols.items() if not v]
    if absent:
        print("Not in export (rules that need them are skipped and the Read Me says so): " + ", ".join(absent))
    live = [r for r in rows if r["status"] == 200]
    os.makedirs(a.out, exist_ok=True)
    out = os.path.join(a.out, "pages.tsv")
    counts = defaultdict(int)
    with open(out, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh, delimiter="\t")
        w.writerow(["id", "path", "words", "index", "title", "h1", "suggested_type", "suggested_node", "type", "node", "city"])
        for i, r in enumerate(live, 1):
            t, node, svc, city = suggest(r, tax, cities)
            counts[t] += 1
            w.writerow([i, r["path"], r["words"], r["index"], r["title"][:90], r["h1"][:90], t, node, t, node, city])
    print(f"Pages: {len(rows)} rows, {len(live)} live (200). Status codes: " +
          ", ".join(f"{k}: {v}" for k, v in sorted(pd.Series([r['status'] for r in rows]).value_counts().items())))
    print("Suggested types: " + ", ".join(f"{k}: {v}" for k, v in sorted(counts.items())))
    print(f"Wrote {out}. Correct the 'type' and 'node' columns (see SKILL.md Step 2); leave suggested_* as they are.")


def overrides_from(a):
    return {k: getattr(a, f"{k}_col", None) for k in COLS}


# ------------------------------------------------------------------ build
def read_pages(path, tax):
    df = pd.read_csv(path, sep="\t", dtype=str, keep_default_na=False, quoting=csv.QUOTE_NONE)
    problems, rows, new_nodes = [], {}, []
    for _, r in df.iterrows():
        t = r["type"].strip().title()
        node = r["node"].strip()
        if t not in PAGE_TYPES:
            problems.append(f"id {r['id']}: type '{r['type']}' is not one of {PAGE_TYPES}")
        n = None
        if node:
            if node.upper().startswith(NEW_PREFIX):
                body = node[len(NEW_PREFIX):]
                if ">" not in body:
                    problems.append(f"id {r['id']}: new node must be 'NEW: Name > Service', got '{node}'")
                else:
                    nm, svc = [x.strip() for x in body.split(">", 1)]
                    parent = tax.lookup(svc)
                    if not parent or parent["level"] != "Service":
                        problems.append(f"id {r['id']}: '{svc}' is not a taxonomy Service")
                    else:
                        n = dict(level="Brand / Device", name=nm, service=parent["service"], category=parent["category"],
                                 subs=[], implies=[], practice=parent["practice"], note="Not in taxonomy", new=True, idx=10**6)
                        new_nodes.append((nm, parent["service"], r["path"]))
            else:
                n = tax.lookup(node)
                if not n:
                    close = difflib.get_close_matches(node.lower(), list(tax.names), n=3)
                    problems.append(f"id {r['id']}: node '{node}' is not in the taxonomy (close: {close}). "
                                    f"Use the exact name, or 'NEW: Name > Service' for a brand/device missing from it.")
        if t == "Service" and not n:
            problems.append(f"id {r['id']}: Service page {r['path']} has no node")
        rows[path_key(r["path"])] = dict(type=t, node=n, city=r.get("city", "").strip())
    if problems:
        print("pages.tsv problems:")
        for p in problems[:60]:
            print("  " + p)
        sys.exit(f"{len(problems)} problem(s) in {path}. Fix them and re-run build.")
    return rows, new_nodes


def tier(v):
    return "No data" if v == 0 else "High" if v >= HIGH else "Medium" if v >= MEDIUM else "Low"


PRIORITY = {  # gap type: (High, Medium, Low / No data)
    "Indexability blocker": ("P1", "P1", "P1"),
    "Empty / thin page": ("P1", "P1", "P1"),
    "Missing page (offered)": ("P1", "P2", "P3"),
    "No Service page (brand pages only)": ("P2", "P3", "P3"),
    "Missing page (blog only)": ("P2", "P3", "P3"),
    "Section, no page": ("P3", "P3", "P3"),
    "Light page": ("P3", "P3", "P3"),
    "Confirm Sub-Service, then page": ("P2", None, None),
    "Ask client: not offered?": ("Ask", None, None),
}


def prio(gap, vol):
    p = PRIORITY[gap]
    return p[0] if vol >= HIGH else p[1] if vol >= MEDIUM else p[2]


def cmd_build(a):
    tax = Taxonomy(a.taxonomy)
    rows, cols = load_crawl(a.crawl, overrides_from(a))
    indexability(rows)
    pages_cfg, new_nodes = read_pages(a.pages, tax)
    cities = cities_list(a.cities)
    practice = a.practice.lower()
    has_h2 = bool(cols["h2"])
    live = [r for r in rows if r["status"] == 200]
    for r in live:
        cfg = pages_cfg.get(r["key"])
        r["type"] = cfg["type"] if cfg else "Utility"
        r["node"] = cfg["node"] if cfg else None
        r["city"] = cfg["city"] if cfg else ""

    # ---- keywords: classify with the taxonomy every run, so a taxonomy edit applies immediately
    kw = tax.kw.copy()
    cls = kw["Keyword"].map(tax.classify_keyword)
    kw[["Class", "Matched Level", "Service", "Sub-Service", "Brand / Device", "Concern"]] = pd.DataFrame(cls.tolist(), index=kw.index)
    kw["kl"] = kw["Keyword"].str.lower()
    city_rx = re.compile("|".join(re.escape(c) for c in cities)) if cities else None
    geo = kw["Geo Modified"].eq("Y")
    own = kw["kl"].map(lambda k: bool(city_rx and city_rx.search(k))) | kw["Place"].str.lower().isin(cities)
    kw["Other Market"] = geo & ~own
    svc_kw = kw[kw["Class"] == "Service"]
    vol_svc = svc_kw.groupby("Service")["Volume"].sum()
    vol_sub = svc_kw[svc_kw["Sub-Service"] != ""].groupby(["Service", "Sub-Service"])["Volume"].sum()
    vol_brand = svc_kw[svc_kw["Brand / Device"] != ""].groupby("Brand / Device")["Volume"].sum()

    def node_vol(n):
        if n["level"] == "Service":
            return int(vol_svc.get(n["service"], 0))
        if n["level"] == "Sub-Service":
            return int(vol_sub.get((n["service"], n["name"]), 0))
        return int(vol_brand.get(n["name"], 0))

    # ---- evidence per node
    def fields(r, which):
        if which == "page":
            return " | ".join([slugtext(r["path"]), r["title"].split("|")[0], r["h1"]])
        if which == "h2":
            return " | ".join([r["h2"], r["h2b"]])
        if which == "meta":
            return r["meta"]
        if which == "all":
            return " | ".join([slugtext(r["path"]), r["title"], r["h1"], r["h2"], r["h2b"], r["meta"]])
        return ""

    in_scope = [n for n in tax.nodes if n["level"] != "Concern" and fits_practice(n["practice"], practice)]
    extra = {}
    for nm, svc, p in new_nodes:
        par = tax.lookup(svc)
        extra[nm] = dict(level="Brand / Device", name=nm, service=par["service"], category=par["category"], subs=[], implies=[],
                         practice=par["practice"], note="Not in taxonomy", new=True, idx=10**6, rx=None)
    nodes = in_scope + list(extra.values())
    ev = {}
    for n in nodes:
        e = dict(page=[], child=[], section=[], meta=[], concern=[], blog=[], gallery=[])
        for r in live:
            rn = r["node"]
            if r["type"] == "Service":
                if rn and rn["name"] == n["name"] and rn["service"] == n["service"]:
                    e["page"].append(r)
                    continue
                if n["level"] == "Service" and rn and rn["service"] == n["service"]:
                    e["child"].append(r)
                    continue
                if n.get("rx") is None:
                    continue
                if has_h2 and n["rx"].search(fields(r, "h2")):
                    e["section"].append(r)
                elif n["rx"].search(fields(r, "meta")):
                    e["meta"].append(r)
            elif n.get("rx") is None:
                continue
            elif r["type"] == "Concern" and n["rx"].search(fields(r, "page") + " | " + fields(r, "h2")):
                e["concern"].append(r)
            elif r["type"] == "Blog" and n["rx"].search(fields(r, "page") + " | " + fields(r, "h2")):
                e["blog"].append(r)
            elif r["type"] == "Gallery" and n["rx"].search(slugtext(r["path"])):
                e["gallery"].append(r)
        ev[(n["level"], n["name"], n["service"])] = e

    # ---- offered status with implication
    reasons = defaultdict(list)

    def mark(level, name, service, why):
        reasons[(level, name, service)].append(why)
    for n in nodes:
        e = ev[(n["level"], n["name"], n["service"])]
        k = (n["level"], n["name"], n["service"])
        if e["page"]:
            mark(*k, "Has page")
        if e["section"]:
            mark(*k, "In H2 on " + e["section"][0]["path"])
        if e["gallery"]:
            mark(*k, f"{len(e['gallery'])} before & after gallery page(s)")
        if e["concern"]:
            mark(*k, "Named on concern page " + e["concern"][0]["path"])
        if e["meta"]:
            mark(*k, "Named in meta description of " + e["meta"][0]["path"])
        if e["blog"]:
            mark(*k, f"{len(e['blog'])} blog post(s)")
        if any(e[x] for x in ("page", "section", "gallery", "concern", "meta", "blog")) and n["level"] != "Service":
            implied = set(n["subs"]) | set(n["implies"]) | {n["service"]}
            for nm in implied:
                tn = tax.lookup(nm)
                if tn and tn["service"] == n["service"] and tn["level"] in ("Service", "Sub-Service"):
                    mark(tn["level"], tn["name"], tn["service"], f"Implied by {n['name']}")
    strong = ("Has page", "In H2", "before & after", "Named on concern page", "Implied by")
    confirmed_services = set()
    status = {}
    for n in nodes:
        k = (n["level"], n["name"], n["service"])
        rs = reasons.get(k, [])
        s = "Confirmed" if any(x.startswith(strong) or "before & after" in x for x in rs) else "Likely" if rs else "No evidence"
        status[k] = s
        if n["level"] == "Service" and s == "Confirmed":
            confirmed_services.add(n["service"])
    for n in nodes:
        k = (n["level"], n["name"], n["service"])
        if n["level"] == "Sub-Service" and status[k] == "No evidence" and n["service"] in confirmed_services:
            status[k] = "Unknown (parent offered)"

    # ---- coverage, gaps
    def best_page(ps):
        return max(ps, key=lambda r: (r["index"] == "Indexable", r["words"] >= THIN, r["words"]))

    out = []
    for n in nodes:
        k = (n["level"], n["name"], n["service"])
        e = ev[k]
        best = best_page(e["page"]) if e["page"] else None
        if best is not None:
            real_noidx = [r for r in e["page"] if r["index"] != "Indexable" and r["words"] >= THIN]
            if real_noidx:
                best = max(real_noidx, key=lambda r: r["words"])
                cov = "Page exists, not indexable"
            elif best["index"] != "Indexable":
                cov = "Page exists, not indexable"
            elif best["words"] < THIN:
                cov = "Empty / thin page"
            elif best["words"] < LIGHT:
                cov = "Light page"
            else:
                cov = "Page"
        elif e["child"]:
            cov = "Only Brand / Sub-Service pages"
        elif e["section"]:
            cov = "Section (H2) only"
        elif e["gallery"]:
            cov = "Gallery only"
        elif e["concern"]:
            cov = "Concern page only"
        elif e["meta"]:
            cov = "Meta description mention only"
        elif e["blog"]:
            cov = "Blog only"
        else:
            cov = "None"
        vol = node_vol(n) if not n.get("new") else 0
        st = status[k]
        gap = ""
        if not flagged_do_not(n["note"]):
            if cov == "Page exists, not indexable":
                gap = "Indexability blocker"
            elif cov == "Empty / thin page":
                gap = "Empty / thin page"
            elif cov == "Light page":
                gap = "Light page"
            elif cov == "Only Brand / Sub-Service pages":
                gap = "No Service page (brand pages only)"
            elif cov == "Section (H2) only":
                gap = "Section, no page"
            elif cov == "Blog only":
                gap = "Missing page (blog only)"
            elif cov in ("Gallery only", "Concern page only", "Meta description mention only") or (cov == "None" and st in ("Confirmed", "Likely")):
                gap = "Missing page (offered)"
            elif st == "Unknown (parent offered)" and vol >= HIGH:
                gap = "Confirm Sub-Service, then page"
            elif (st == "No evidence" and n["level"] == "Service" and vol >= HIGH
                  and n["category"] != "Beauty Services"):
                gap = "Ask client: not offered?"
        p = prio(gap, vol) if gap else ""
        if gap and p is None:
            gap, p = "", ""
        out.append(dict(node=n, key=k, e=e, best=best, cov=cov, vol=vol, status=st, reasons=reasons.get(k, []),
                        gap=gap, prio=p))

    # ---- keywords per recommended page
    def page_kws(n):
        if n.get("new"):
            return svc_kw.iloc[0:0]
        if n["level"] == "Service":
            m = (svc_kw["Service"] == n["service"]) & (svc_kw["Sub-Service"] == "") & (svc_kw["Brand / Device"] == "")
        elif n["level"] == "Sub-Service":
            m = (svc_kw["Service"] == n["service"]) & (svc_kw["Sub-Service"] == n["name"])
        else:
            m = (svc_kw["Brand / Device"] == n["name"]) & (svc_kw["Sub-Service"] == "")
        return svc_kw[m].sort_values("Volume", ascending=False)

    gsc = load_gsc(a.gsc, a.gsc_tab) if a.gsc else None

    def display(n):
        return n["name"] if n["level"] == "Service" else f"{n['name']} ({n['service']})"

    recs, kmap, excl = [], [], []
    for o in out:
        n, gap = o["node"], o["gap"]
        if not gap:
            continue
        if n["level"] == "Brand / Device" and gap not in ("Indexability blocker", "Empty / thin page", "Light page"):
            continue
        e = o["e"]
        label, url, existing = "", "", ""
        if gap == "Indexability blocker":
            label, url = LABELS[0], o["best"]["path"]
        elif gap == "Empty / thin page":
            label, url = LABELS[1], o["best"]["path"]
        elif gap == "Light page":
            label, url = LABELS[8], o["best"]["path"]
        elif gap == "No Service page (brand pages only)":
            kids = e["child"]
            label = LABELS[6] if all(r["node"]["level"] == "Sub-Service" for r in kids) else LABELS[5]
            existing = "; ".join(f"{r['node']['name']}: {r['path']}" for r in kids)
        elif gap == "Missing page (blog only)":
            label = LABELS[7]
        elif gap == "Missing page (offered)":
            label = LABELS[2] if e["gallery"] else LABELS[3] if e["concern"] else LABELS[4]
        elif gap == "Section, no page":
            label = LABELS[4]
        elif gap == "Confirm Sub-Service, then page":
            label = LABELS[9]
        elif gap == "Ask client: not offered?":
            label = LABELS[10]
        k = page_kws(n)
        own = k[~k["Other Market"]]
        name = display(n)
        for x in own.itertuples():
            kmap.append([name, url, x.Keyword, int(x.Volume)])
        for x in k[k["Other Market"]].itertuples():
            excl.append([x.Keyword, int(x.Volume), name, f"Other market (Geo Modified: {x.Place})", ""])
        kw_cell = ", ".join(own["Keyword"]) if len(own) else NO_KW_TEXT
        gsc_impr = None
        if gsc is not None:
            qs = set(own["kl"])
            gsc_impr = sum(v[0] for page in gsc.values() for q, v in page.items() if q in qs)
        o["rec"] = dict(name=name, url=url, existing=existing, label=label, kw=kw_cell, vol=int(own["Volume"].sum()), gsc=gsc_impr)
        recs.append(o)
    recs.sort(key=lambda o: (LABELS.index(o["rec"]["label"]), -o["rec"]["vol"]))

    # ---- page briefs and recommendations text
    for o in out:
        o["recommendation"] = recommendation(o, out, tax, svc_kw, ev)

    # ---- workbook
    write_workbook(a, tax, cols, rows, live, out, recs, kmap, excl, gsc is not None, has_h2, cities, practice, new_nodes)

    # ---- QA
    qa(a.out, recs, kmap, live, out)
    report(out, recs, live, new_nodes, kw)


def recommendation(o, out, tax, svc_kw, ev):
    g, n = o["gap"], o["node"]
    if not g:
        return ""
    if g == "Empty / thin page" and n["level"] == "Service":
        return hub_brief(o, out, svc_kw)
    one = {
        "Indexability blocker": "Make the page indexable (canonical, noindex or redirect). The content already exists.",
        "Empty / thin page": f"Build out the {n['name']} page. It exists but has almost no content.",
        "Light page": f"Expand toward the depth of the site's main service pages: areas, candidacy, results, cost, FAQs.",
        "No Service page (brand pages only)": f"Build a {n['name']} Service page that links to the existing brand / Sub-Service page(s), or retarget a single brand page to the Service term.",
        "Missing page (blog only)": f"Only blog posts cover {n['name']}. Create a Service page and link the posts to it.",
        "Missing page (offered)": f"Create a {n['name']} page." + (" Before & after galleries already exist." if o["e"]["gallery"] else
                                                                    " Named on a concern page." if o["e"]["concern"] else
                                                                    " Confirm with client first (only mentioned)." if o["status"] == "Likely" else ""),
        "Section, no page": "Covered as an H2 section. Separate page only if volume justifies it.",
        "Confirm Sub-Service, then page": f"Ask the client if they offer {n['name']}. If yes, create a page (or at minimum an H2 section on the parent).",
        "Ask client: not offered?": "No sign the practice offers this. Ask the client.",
    }
    return one[g]


def hub_brief(o, out, svc_kw):
    n, e = o["node"], o["e"]
    parts = []
    real = [r for r in e["page"] if r["words"] >= THIN]
    thin = [r for r in e["page"] if r["words"] < THIN]
    main = o["best"]
    parts.append(f"Build {main['path']} into the main {n['name']} page. It has {main['words']} content words today.")
    for r in thin:
        if r is main:
            continue
        parts.append(f"{r['path']} ({r['words']} words, H1 \"{r['h1']}\") duplicates it: merge and 301 it.")
    by_svc = [x for x in out if x["node"]["service"] == n["service"]]
    subs = [x["node"]["name"] for x in by_svc if x["node"]["level"] == "Sub-Service" and x["status"] in ("Confirmed", "Likely")]
    ask = [x["node"]["name"] for x in by_svc if x["node"]["level"] == "Sub-Service" and x["status"].startswith("Unknown") and x["vol"] >= HIGH]
    brands = [x["node"]["name"] for x in by_svc if x["node"]["level"] == "Brand / Device" and x["best"] is not None]
    if subs:
        parts.append("Add a section for each offered Sub-Service, linking to its own page once built: " + ", ".join(subs) + ".")
    if ask:
        parts.append("Sub-Services to confirm with the client: " + ", ".join(ask) + ".")
    if brands:
        parts.append("Compare and link the brand pages: " + ", ".join(brands) + ".")
    return " ".join(parts)


# ------------------------------------------------------------------ gsc (same sheet format as service-page-title-h1-review)
GSC_TAB = "Claude Use This"
GSC_COLS = {"page": ("page", "landing page", "url", "address", "top pages"),
            "query": ("query", "top queries", "search query"),
            "clicks": ("clicks",), "impressions": ("impressions",)}


def gsc_tabs(path):
    raw = {}
    if path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        for name, df in pd.read_excel(path, sheet_name=None, dtype=str, header=None).items():
            raw[name] = df.fillna("").values.tolist()
    else:
        name = "(file)"
        with open(path, encoding="utf-8-sig", newline="") as fh:
            for line in fh:
                if line.startswith("## Sheet name:"):
                    name = line.split(":", 1)[1].strip()
                    raw[name] = []
                elif line.strip():
                    raw.setdefault(name, []).append(next(csv.reader([line])))
    tabs = {}
    for name, rows in raw.items():
        for hi, row in enumerate(rows[:10]):
            cells = [str(c).strip().lower() for c in row]
            idx = {k: next((cells.index(x) for x in names if x in cells), None) for k, names in GSC_COLS.items()}
            if all(v is not None for v in idx.values()):
                got = []
                for r in rows[hi + 1:]:
                    if len(r) <= max(idx.values()):
                        continue
                    try:
                        c = float(str(r[idx["clicks"]]).replace(",", "") or 0)
                        i = float(str(r[idx["impressions"]]).replace(",", "") or 0)
                    except ValueError:
                        continue
                    if str(r[idx["query"]]).strip():
                        got.append((str(r[idx["page"]]).strip(), str(r[idx["query"]]).strip(), c, i))
                tabs[name] = got
                break
    return tabs


def load_gsc(path, tab=None):
    if not os.path.isfile(path):
        sys.exit(f"GSC file not found: {path}")
    tabs = gsc_tabs(path)
    if not tabs:
        sys.exit(f"{path}: no tab has Page, Query, Clicks and Impressions columns, so it is not Search Console data "
                 "(it may be the taxonomy or another sheet). Stop and tell the person.")
    if not tab:
        tab = next((n for n in tabs if n.strip().lower() == GSC_TAB.lower()), None)
    if tab:
        if tab not in tabs:
            sys.exit(f"--gsc-tab '{tab}' is not a tab with those columns. Tabs: {', '.join(tabs)}")
        use = tab
    elif len(tabs) == 1:
        use = next(iter(tabs))
    else:
        listing = "; ".join(f"'{n}' ({len(r)} rows)" for n, r in tabs.items())
        sys.exit(f"{path} has several tabs with GSC columns: {listing}. Ask the person which tab to use, then re-run "
                 f"with --gsc-tab NAME (or name the tab \"{GSC_TAB}\"). Never add the tabs together.")
    total = {}
    for page, q, c, i in tabs[use]:
        v = total.setdefault(path_key(page), {}).setdefault(q.lower(), [0, 0])
        v[0] += int(i)
        v[1] += int(c)
    print(f"  GSC tab '{use}': {len(tabs[use])} query rows, {len(total)} pages")
    return total


# ------------------------------------------------------------------ workbook
def add_sheet(wb, title, headers, rows, widths, numeric=(), yellow=(), first_row=1, name=None):
    ws = wb.create_sheet(title)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(first_row, c, h)
        cell.font, cell.fill = HEAD, HEAD_FILL
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for i, r in enumerate(rows, first_row + 1):
        for c, v in enumerate(r, 1):
            cell = ws.cell(i, c, v)
            cell.font = ARIAL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if headers[c - 1] in numeric:
                cell.number_format = "#,##0"
            if headers[c - 1] in yellow:
                cell.fill = YELLOW
    ws.freeze_panes = ws.cell(first_row + 1, 1)
    t = Table(displayName=name or re.sub(r"\W", "", title), ref=f"A{first_row}:{get_column_letter(len(headers))}{max(len(rows), 1) + first_row}")
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(t)
    return ws


def write_workbook(a, tax, cols, rows, live, out, recs, kmap, excl, has_gsc, has_h2, cities, practice, new_nodes):
    wb = Workbook()
    rm = wb.active
    rm.title = "Read Me"
    rm.column_dimensions["A"].width, rm.column_dimensions["B"].width = 26, 100
    rm["A1"] = f"Service Gap Audit: {a.domain or os.path.basename(a.out).split('_service_gap')[0]}"
    rm["A1"].font = Font(name="Arial", bold=True, size=14)
    lines = [
        ("Inputs", f"Crawl export: {os.path.basename(a.crawl)} ({len(rows)} URLs, {len(live)} live). Taxonomy: {tax.version}. "
                   f"Search Console: {'used (' + os.path.basename(a.gsc) + ')' if has_gsc else 'not provided, so there is no Client GSC Impressions column. Provide the GSC Google Sheet and re-run to add it.'}"),
        ("Fields used", ", ".join(k for k, v in cols.items() if v) + (". No H2 columns in the export, so H2 section matching was skipped." if not has_h2 else ".")),
        ("Practice type", f"{practice.title()}. Service-area cities kept: {', '.join(c.title() for c in cities) or 'none given'}."),
        ("Level names", "Category > Service > Sub-Service > Brand / Device. A Concern is a searched problem, never a page gap on its own."),
        ("Page types", "Set per page by reviewing the crawl (pages.tsv), not from folder names: Service, Concern, Gallery, Blog, Location, Utility."),
        ("Offered Status", "Confirmed = own page, H2 section, gallery, concern page, or implied by a brand / Sub-Service. Likely = only a meta description or blog mention. Unknown (parent offered) = Sub-Service of a Confirmed Service. No evidence."),
        ("Priority", "P1 fix first, P2 next, P3 later, Ask = confirm with the client. Set by gap type and volume tier (High >= 5,000, Medium 2,000-4,999, Low < 2,000)."),
        ("Competitor Volume", "Summed from the taxonomy's competitor keyword list (US national). Ranks gaps against each other; not local demand."),
        ("Limits", "The export has no body text: a service mentioned only in body copy shows as not found. Check the parent page before calling it missing."),
        ("Thresholds", f"Thin < {THIN} content words, light < {LIGHT}."),
        ("Search volume", "Relevant Keywords = keywords for that page, counted once on the most specific page. Keywords tied to another market (Geo Modified) are excluded unless they name a service-area city; see Excluded Keywords."),
        ("Not in taxonomy", "; ".join(f"{nm} ({svc}): {p}" for nm, svc, p in new_nodes) or "None"),
        ("Yellow columns", "Fill in: Keep? on Excluded Keywords, Client Confirmed and Notes on Gap List."),
    ]
    for i, (k, v) in enumerate(lines, 3):
        rm.cell(i, 1, k).font = Font(name="Arial", bold=True)
        c = rm.cell(i, 2, v)
        c.font, c.alignment = ARIAL, Alignment(wrap_text=True, vertical="top")

    # Recommendations
    h = ["Service Page Recommended", "Current URL", "Existing Brand / Sub-Service Page(s)", "Why Recommended",
         "Relevant Keywords", "Total National Search Volume"] + (["Client GSC Impressions"] if has_gsc else [])
    rws = [[o["rec"]["name"], o["rec"]["url"], o["rec"]["existing"], o["rec"]["label"], o["rec"]["kw"], o["rec"]["vol"]]
           + ([o["rec"]["gsc"]] if has_gsc else []) for o in recs]
    ws = add_sheet(wb, "Recommendations", h, rws, [36, 46, 46, 38, 90, 14] + ([14] if has_gsc else []),
                   numeric=("Total National Search Volume", "Client GSC Impressions"), first_row=2, name="Recs")
    ws.merge_cells("A1:D1")
    ws["A1"] = "Page"
    ws.merge_cells(f"E1:{get_column_letter(len(h))}1")
    ws["E1"] = "Search Volume"
    for c, fill in (("A1", "D9D9D9"), ("E1", "BDD7EE")):
        ws[c].font, ws[c].fill = Font(name="Arial", bold=True), PatternFill("solid", fgColor=fill)

    add_sheet(wb, "Keyword Map", ["Service Page Recommended", "Current URL", "Keyword", "National Search Volume"], kmap,
              [36, 46, 48, 14], numeric=("National Search Volume",), name="KeywordMap")
    add_sheet(wb, "Excluded Keywords", ["Keyword", "National Search Volume", "Service Page Recommended", "Exclusion Reason", "Keep? (Y to add back)"],
              sorted(excl, key=lambda r: -r[1]), [44, 14, 36, 36, 14], numeric=("National Search Volume",),
              yellow=("Keep? (Y to add back)",), name="Excluded")

    gl = [o for o in out if o["gap"]]
    gl.sort(key=lambda o: (["P1", "P2", "P3", "Ask"].index(o["prio"]), o["gap"], -o["vol"]))
    gh = ["Priority", "Gap Type", "Category", "Service", "Level", "Service / Sub-Service / Brand", "Offered Status", "Offered Evidence",
          "Coverage", "Best Page", "Content Words", "Indexability", "All Pages for Node", "Competitor Volume", "Volume Tier",
          "Recommendation", "Client Confirmed", "Notes"]
    grows = [[o["prio"], o["gap"], o["node"]["category"], o["node"]["service"], o["node"]["level"], o["node"]["name"], o["status"],
              "; ".join(dict.fromkeys(o["reasons"])), o["cov"], o["best"]["path"] if o["best"] else "",
              o["best"]["words"] if o["best"] else None, o["best"]["index"] if o["best"] else "",
              "; ".join(f"{r['path']} ({r['words']}w, {r['index']})" for r in o["e"]["page"]), o["vol"], tier(o["vol"]),
              o["recommendation"], "", ""] for o in gl]
    ws = add_sheet(wb, "Gap List", gh, grows, [8, 24, 18, 22, 12, 26, 14, 36, 20, 38, 10, 20, 44, 12, 10, 70, 12, 30],
                   numeric=("Content Words", "Competitor Volume"), yellow=("Client Confirmed", "Notes"), name="Gaps")
    for i, o in enumerate(gl, 2):
        ws.cell(i, 1).fill = PatternFill("solid", fgColor=PRIO_FILL[o["prio"]])
    dv = DataValidation(type="list", formula1='"Yes,No,Planned"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"Q2:Q{max(len(gl), 1) + 1}")

    # Page Inventory
    same = defaultdict(list)
    for r in live:
        if r["type"] == "Service" and r["node"]:
            same[(r["node"]["service"], r["node"]["name"])].append(r)
    inv = []
    order = {t: i for i, t in enumerate(["Service", "Concern", "Location", "Blog", "Utility"])}
    for r in sorted([r for r in live if r["type"] != "Gallery"], key=lambda r: (order.get(r["type"], 9), r["path"])):
        n = r["node"]
        issues = []
        if r["index"] != "Indexable":
            issues.append(r["index"])
        if r["type"] == "Service" and r["words"] < THIN:
            issues.append(f"Empty / thin ({r['words']} content words)")
        elif r["type"] == "Service" and r["words"] < LIGHT:
            issues.append(f"Light ({r['words']} content words)")
        if not r["h1"]:
            issues.append("Missing H1")
        dup = ""
        if n and len(same[(n["service"], n["name"])]) > 1:
            others = [x for x in same[(n["service"], n["name"])] if x is not r]
            dup = "; ".join(f"{x['path']} ({x['words']}w)" for x in others)
            if r["words"] < THIN and any(x["words"] >= THIN for x in others):
                issues.append("Thin copy of another page for the same node" + (f" ({r['city'].title()})" if r["city"] else ""))
        inv.append([r["path"], r["type"], n["level"] if n else "", n["service"] if n else "", n["name"] if n else "",
                    r["city"].title(), r["words"], r["index"], "Y" if not r["h1"] else "", dup, "; ".join(issues),
                    r["title"], r["h1"], r["h2"], r["h2b"], r["meta"]])
    add_sheet(wb, "Page Inventory", ["Path", "Page Type", "Primary Level", "Primary Service", "Primary Service / Sub-Service / Brand",
                                     "City", "Content Words", "Indexability", "H1 Missing", "Duplicate Of", "Page Issues", "Title", "H1",
                                     "H2", "H2 (second)", "Meta Description"], inv,
              [48, 10, 12, 22, 26, 12, 10, 20, 8, 40, 40, 40, 30, 30, 30, 50], numeric=("Content Words",), name="Pages")

    cv = sorted(out, key=lambda o: (o["node"]["category"], o["node"]["service"], ["Service", "Sub-Service", "Brand / Device"].index(o["node"]["level"]), -o["vol"]))
    add_sheet(wb, "Service Coverage", ["Category", "Service", "Level", "Service / Sub-Service / Brand", "Offered Status", "Coverage",
                                       "Best Page", "Content Words", "Indexability", "Competitor Volume", "Gap Type", "Status Note"],
              [[o["node"]["category"], o["node"]["service"], o["node"]["level"], o["node"]["name"], o["status"], o["cov"],
                o["best"]["path"] if o["best"] else "", o["best"]["words"] if o["best"] else None,
                o["best"]["index"] if o["best"] else "", o["vol"], o["gap"], o["node"]["note"]] for o in cv],
              [18, 24, 12, 30, 16, 26, 40, 10, 20, 12, 26, 36], numeric=("Content Words", "Competitor Volume"), name="Coverage")

    gal = defaultdict(list)
    for r in live:
        if r["type"] == "Gallery":
            slug = re.sub(r"[-_]?\d+$", "", r["path"].rstrip("/").split("/")[-1] or "/")
            gal[slug].append(r)
    grows = []
    for slug, rs in sorted(gal.items()):
        n = tax.page_node(slugtext(slug))
        grows.append([slug, len(rs), f"{n['service']} > {n['name']}" if n else "Package name? Ask client"])
    add_sheet(wb, "Gallery Evidence", ["Slug", "Gallery Pages", "Maps To"], grows, [30, 12, 40], name="Gallery")
    wb.save(a.out)
    print(f"Saved {a.out}")


# ------------------------------------------------------------------ QA and report
def qa(path, recs, kmap, live, out):
    checks = []
    checks.append(("Every Why Recommended value is one of the labels", all(o["rec"]["label"] in LABELS for o in recs)))
    sums = defaultdict(int)
    for name, url, k, v in kmap:
        sums[name] += v
    checks.append(("Total National Search Volume equals the Keyword Map rows", all(sums.get(o["rec"]["name"], 0) == o["rec"]["vol"] for o in recs)))
    ok = True
    for o in recs:
        if o["rec"]["label"] in (LABELS[5], LABELS[6]):
            kids = o["e"]["child"]
            ok &= bool(o["rec"]["existing"]) and all(r["node"]["service"] == o["node"]["service"] for r in kids)
    checks.append(("Brand-only rows name their existing page(s) under the same Service", ok))
    checks.append(("No gap on a node marked Do not flag / NOT US-approved", not any(o["gap"] and flagged_do_not(o["node"]["note"]) for o in out)))
    checks.append(("Every Service page has a node", all(r["node"] for r in live if r["type"] == "Service")))
    checks.append(("Brand / Device rows are named with their parent", all("(" in o["rec"]["name"] for o in recs if o["node"]["level"] != "Service")))
    passed = True
    for name, ok in checks:
        print(("PASS  " if ok else "FAIL  ") + name)
        passed &= ok
    if not passed:
        print("Fix every FAIL and re-run build.")


def report(out, recs, live, new_nodes, kw):
    c = defaultdict(int)
    for o in recs:
        c[o["rec"]["label"]] += 1
    print(f"Recommendations: {len(recs)} rows")
    for l in LABELS:
        if c[l]:
            print(f"  {c[l]:>3}  {l}")
    top = [o for o in recs if o["rec"]["vol"]][:10]
    print("Top rows by label order (check their Relevant Keywords for other cities or wrong services):")
    for o in top:
        print(f"  {o['rec']['name']} | {o['rec']['label']} | {o['rec']['vol']:,} | {o['rec']['kw'][:110]}")
    if new_nodes:
        print("Not in taxonomy (list for the taxonomy pull request): " + "; ".join(f"{a} > {b}" for a, b, _ in new_nodes))


# ------------------------------------------------------------------ main
def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("inspect", "build"):
        sp = sub.add_parser(name)
        sp.add_argument("--crawl", required=True)
        sp.add_argument("--taxonomy", default=DEFAULT_TAXONOMY)
        sp.add_argument("--cities", default="", help="Service-area cities, comma separated")
        sp.add_argument("--practice", default="both", choices=["medspa", "surgical", "both"])
        sp.add_argument("--out", required=True)
        for k in COLS:
            sp.add_argument(f"--{k}-col")
        if name == "build":
            sp.add_argument("--pages", required=True)
            sp.add_argument("--gsc")
            sp.add_argument("--gsc-tab")
            sp.add_argument("--domain", default="")
    a = ap.parse_args()
    (cmd_inspect if a.cmd == "inspect" else cmd_build)(a)


if __name__ == "__main__":
    main()
