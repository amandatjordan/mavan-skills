#!/usr/bin/env python3
"""Service Page H1 & Title Tag Review.

  python build.py inspect --crawl export.csv --out work
  python build.py build   --crawl export.csv --keywords work/keywords.tsv --out review.xlsx
                          [--gsc GSC_SHEET [--gsc-tab TAB]] [--brand "term, term"] [--cities "City, City"]

inspect: shows which export columns are used, writes work/pages.tsv (one compact line per page, for the AI to read), and lists pages that look like
         they are not service pages (tags, categories, blog posts).
build:   takes the AI's topic and target keywords for every page, adds each page's top Search Console queries
         when --gsc is given, flags the pages a person should review and builds the two-tab workbook.
         --gsc reads the provided Google Sheet of page + query rows. Queries with fewer than 5 impressions and
         queries containing a --brand term are dropped.
The AI decides what each page is about and its keywords. It does not rank or judge Search Console queries.
Every number, count and join happens here. Rules: references/rules.md.
"""
import argparse
import csv
import os
import re
import sys
import unicodedata
from collections import defaultdict
from urllib.parse import unquote, urlsplit

import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

URL_NAMES = ["address", "url", "page url"]
H1_NAMES = ["h1-1", "h1", "h1 1"]
TITLE_NAMES = ["title 1", "title", "page title"]
MAX_KEYWORDS = 8   # the AI proposes at most this many per page
MIN_IMPR, TOP_QUERIES = 5, 10   # a query needs this many impressions; this many per page
NON_SERVICE = re.compile(r"/(tags?|category|categories|page-category|author|blog|news|feed)(/|$)|/page/\d+|/\d{4}/\d{2}/", re.I)
NO_GSC_TEXT = "No GSC data for this page (or fewer than 5 impressions per query)"
STOP = {"and", "or", "in", "for", "of", "the", "a", "an", "to", "with", "near", "me", "by", "vs"}


# ---------------------------------------------------------------- crawl / urls
def read_table(path):
    if path.lower().endswith((".xlsx", ".xlsm")):
        return pd.read_excel(path, dtype=str).fillna("")
    return pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig", encoding_errors="replace")


def find_col(df, override, names, label):
    lookup = {c.strip().lower(): c for c in df.columns}
    if override:
        if override.strip().lower() not in lookup:
            sys.exit(f"--{label}-col '{override}' is not a column. Columns: {list(df.columns)}")
        return lookup[override.strip().lower()]
    for n in names:
        if n in lookup:
            return lookup[n]
    sys.exit(f"Could not find the {label} column. Columns: {list(df.columns)}. Re-run with --{label}-col NAME.")


def norm_url(u):
    u = u.strip()
    if not re.match(r"^[a-z][a-z0-9+.-]*://", u, re.I):
        u = "//" + u
    p = urlsplit(u)
    path = unquote(p.path or "/")
    if len(path) > 1:
        path = path.rstrip("/")
    return (p.netloc.lower() + path).lower()


def page_path(u):
    p = urlsplit(u.strip())
    return unquote(p.path) or "/"


def load_crawl(args):
    df = read_table(args.crawl)
    cu = find_col(df, args.url_col, URL_NAMES, "url")
    ch = find_col(df, args.h1_col, H1_NAMES, "h1")
    ct = find_col(df, args.title_col, TITLE_NAMES, "title")
    pages, seen, dups = [], set(), 0
    for _, r in df.iterrows():
        url = r[cu].strip()
        if not url:
            continue
        key = norm_url(url)
        if key in seen:
            dups += 1
            continue
        seen.add(key)
        pages.append({"url": url, "key": key, "path": page_path(url), "h1": r[ch].strip(), "title": r[ct].strip()})
    return pages, (cu, ch, ct), dups


# ------------------------------------------------------------------------ gsc

GSC_TAB = "Claude Use This"   # the tab of the Search Console template that the script uses when it exists
GSC_COLS = {"page": ("page", "landing page", "url", "address", "top pages"),
            "query": ("query", "top queries", "search query"),
            "clicks": ("clicks",), "impressions": ("impressions",)}


def path_key(u):
    p = unquote(urlsplit(u.strip()).path or "/").lower()
    return p.rstrip("/") or "/"


def to_num(x):
    try:
        return float(str(x).replace(",", "").strip() or 0)
    except ValueError:
        return None


def gsc_tabs(path):
    """-> {tab name: [(page, query, clicks, impressions)]} for every tab of the Google Sheet export that has
    Page, Query, Clicks and Impressions columns. Reads an .xlsx download, or the text export where each tab
    starts with a '## Sheet name: NAME' line followed by CSV."""
    raw = {}
    if path.lower().endswith((".xlsx", ".xlsm", ".xls")):
        for name, df in pd.read_excel(path, sheet_name=None, dtype=str, header=None).items():
            raw[name] = df.fillna("").values.tolist()
    else:
        name = "(file)"
        with open(path, encoding="utf-8-sig", newline="") as fh:
            for line in fh:
                if line.startswith("## Sheet name:"):
                    name = line.split(":", 1)[1].strip()
                    raw[name] = []
                elif line.strip():
                    raw.setdefault(name, []).append(next(csv.reader([line])))
    tabs = {}
    for name, rows in raw.items():
        for hi, row in enumerate(rows[:10]):
            cells = [str(c).strip().lower() for c in row]
            idx = {k: next((cells.index(n) for n in names if n in cells), None) for k, names in GSC_COLS.items()}
            if all(v is not None for v in idx.values()):
                out = []
                for r in rows[hi + 1:]:
                    if len(r) <= max(idx.values()):
                        continue
                    c, i = to_num(r[idx["clicks"]]), to_num(r[idx["impressions"]])
                    if c is None or i is None or not str(r[idx["query"]]).strip():
                        continue
                    out.append((str(r[idx["page"]]).strip(), str(r[idx["query"]]).strip(), c, i))
                tabs[name] = out
                break
    return tabs


def load_gsc(path, tab=None):
    """-> {page path: {query: [impressions, clicks]}}. The page's query string (utm and similar) is dropped, so
    the same page under several URLs adds up."""
    if not os.path.isfile(path):
        sys.exit(f"GSC file not found: {path}")
    tabs = gsc_tabs(path)
    if not tabs:
        sys.exit(f"{path}: no tab has Page, Query, Clicks and Impressions columns. Stop and tell the person.")
    if not tab:
        tab = next((n for n in tabs if n.strip().lower() == GSC_TAB.lower()), None)
    if tab:
        if tab not in tabs:
            sys.exit(f"--gsc-tab '{tab}' is not a tab with those columns. Tabs: {{ {', '.join(tabs)} }}")
        use = tab
    elif len(tabs) == 1:
        use = next(iter(tabs))
    else:
        listing = "; ".join(f"'{n}' ({len(r)} rows)" for n, r in tabs.items())
        sys.exit(f"{path} has several tabs with GSC columns: {listing}. Ask the person which tab to use, "
                 "then re-run with --gsc-tab NAME (or name the tab \"" + GSC_TAB + "\"). Never add the tabs together.")
    total = {}
    for page, q, c, i in tabs[use]:
        v = total.setdefault(path_key(page), {}).setdefault(q.lower(), [0, 0])
        v[0] += int(i)
        v[1] += int(c)
    print(f"  GSC tab '{use}': {len(tabs[use])} query rows, {len(total)} pages")
    return total


def top_queries(queries, brand_terms):
    """[(query, impressions, clicks)]: at least MIN_IMPR impressions, no brand term, most impressions first."""
    keep = [(q, i, c) for q, (i, c) in queries.items()
            if i >= MIN_IMPR and not any(b in q for b in brand_terms)]
    return sorted(keep, key=lambda x: (-x[1], -x[2], x[0]))[:TOP_QUERIES]


def dedupe_key(term):
    k = re.sub(r"[\s\-]+", "", term.lower())
    return k[:-1] if len(k) > 3 and k.endswith("s") else k



def tokens(s):
    s = unicodedata.normalize("NFKD", s).encode("ascii", "ignore").decode().lower()
    return re.findall(r"[a-z0-9]+", s)

def term_key(term):
    """Same folding as dedupe_key, so 'med spa', 'medspa' and 'med-spa' (and a plural) are one term."""
    return dedupe_key("".join(tokens(term)))

def phrase_keys(text, maxn=6):
    """Every run of up to maxn consecutive words in the text, as term keys. A term is 'in' the text if its key is here."""
    t = tokens(text)
    return {dedupe_key("".join(t[i:j])) for i in range(len(t)) for j in range(i + 1, min(len(t), i + maxn) + 1)}


def term_tokens(s):
    """Words of a term with plurals folded, so 'injectables' and 'injectable' are the same word."""
    out = []
    for t in tokens(s):
        if t in STOP:
            continue
        if len(t) > 3 and t.endswith("s") and not t.endswith("ss"):
            t = t[:-1]
        out.append(t)
    return out


def read_keywords(path, pages):
    """work/keywords.tsv: id <TAB> topic <TAB> keyword; keyword; ... (most important first)."""
    out, ids = {}, {p["id"] for p in pages}
    with open(path, encoding="utf-8-sig") as fh:
        for line in fh:
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            parts = [x.strip() for x in line.rstrip("\n").split("\t")]
            if not parts[0].isdigit():
                continue  # header row
            i = int(parts[0])
            if i not in ids:
                sys.exit(f"{path}: id {i} is not a page id in work/pages.tsv")
            topic = parts[1] if len(parts) > 1 else ""
            seen, kws = set(), []
            for k in re.split(r"[;|]", parts[2] if len(parts) > 2 else ""):
                k = k.strip()
                if k and dedupe_key(k) not in seen:
                    seen.add(dedupe_key(k))
                    kws.append(k)
            out[i] = (topic, kws[:MAX_KEYWORDS])
    missing = sorted(i for i in ids if not out.get(i, ("", []))[0])
    if missing:
        sys.exit(f"{path}: no topic for page ids {missing[:20]}{' ...' if len(missing) > 20 else ''}. Every page needs one.")
    return out


# ------------------------------------------------------------------- workbook

ARIAL = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="1F3864")
YELLOW = PatternFill("solid", fgColor="FFF2CC")


def add_sheet(wb, first, title, headers, rows, widths, numeric=(), yellow=(), name=None):
    ws = wb.active if first else wb.create_sheet()
    ws.title = title
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(1, c)
        cell.font, cell.fill = HEAD, (YELLOW if h in yellow else HEAD_FILL)
        if h in yellow:
            cell.font = Font(name="Arial", size=10, bold=True)
        cell.alignment = Alignment(vertical="top", wrap_text=True)
        ws.column_dimensions[get_column_letter(c)].width = widths[c - 1]
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = ARIAL
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if headers[cell.column - 1] in numeric:
                cell.number_format = "#,##0"
            if headers[cell.column - 1] in yellow:
                cell.fill = YELLOW
    ws.freeze_panes = "A2"
    ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    t = Table(displayName=name or title.replace(" ", ""), ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    ws.add_table(t)
    return ws


# ----------------------------------------------------------------------- main

# ----------------------------------------------------------------------- main

def short(s, n=80):
    s = " ".join(s.split())
    return s if len(s) <= n else s[:n - 1] + "…"


def cmd_inspect(args):
    pages, cols, dups = load_crawl(args)
    os.makedirs(args.out, exist_ok=True)
    odd = [p for p in pages if NON_SERVICE.search(p["path"])]
    with open(os.path.join(args.out, "pages.tsv"), "w", encoding="utf-8") as fh:
        fh.write("id\tpath\th1\ttitle\tnote\n")
        for i, p in enumerate(pages, 1):
            fh.write(f"{i}\t{p['path']}\t{short(p['h1'])}\t{short(p['title'].split('|')[0])}\t"
                     f"{'likely not a service page' if p in odd else ''}\n")
    print(f"Columns used: URL='{cols[0]}', H1='{cols[1]}', Title='{cols[2]}'")
    print(f"Pages: {len(pages)} (duplicate URLs skipped: {dups}); blank H1: {sum(not p['h1'] for p in pages)}; "
          f"blank title: {sum(not p['title'] for p in pages)}")
    if odd:
        print(f"Likely not service pages ({len(odd)}); kept and flagged, ask the person once whether to drop them:")
        for p in odd:
            print("  " + p["path"])
    print(f"Wrote {os.path.join(args.out, 'pages.tsv')}")


def cmd_build(args):
    pages, cols, dups = load_crawl(args)
    for i, p in enumerate(pages, 1):
        p["id"] = i
    kw = read_keywords(args.keywords, pages)
    crawl_keys = {p["key"] for p in pages}
    gsc = load_gsc(args.gsc, args.gsc_tab) if args.gsc else None
    brand = [b.strip().lower() for b in (args.brand or "").split(",") if b.strip()]
    cities = [c.strip() for c in (args.cities or "").split(",") if c.strip()]
    city_keys = {term_key(c) for c in cities}
    print(f"Pages: {len(pages)} (duplicates skipped {dups})")
    print("GSC: " + ("used; brand terms dropped: " + (", ".join(brand) or "none given") if gsc is not None
                     else "not used (no --gsc): no Top GSC Queries column"))
    print("City check: " + (", ".join(cities) if cities else "skipped (no --cities)"))

    with_data = 0
    for p in pages:
        p["topic"], kws = kw[p["id"]]
        p["notes"], p["flags"] = [], []
        p["keywords"] = kws[:5]
        p["queries"] = []
        if gsc is not None:
            p["queries"] = top_queries(gsc.get(path_key(p["url"]), {}), brand)
            with_data += bool(p["queries"])
        if NON_SERVICE.search(p["path"]):
            p["flags"].append(("scope", "Likely not a service page (tag, category or blog path)"))
        if not p["keywords"]:
            p["flags"].append(("nokeywords", "No target keywords. Choose them by hand"))
        if not p["h1"]:
            p["flags"].append(("blankh1", "Blank H1"))
        if not p["title"]:
            p["flags"].append(("blanktitle", "Blank Title Tag"))
        have = phrase_keys(p["title"]) | phrase_keys(p["h1"])
        if p["keywords"] and not any(term_key(k) in have for k in p["keywords"]):
            p["flags"].append(("notrep", "None of the top keywords appear in the Title Tag or H1"))
        if city_keys:
            for label, text in (("Title Tag", p["title"]), ("H1", p["h1"])):
                if text and not (city_keys & phrase_keys(text)):
                    p["flags"].append(("nocity", f"No location in {label}"))

    if gsc is not None:
        rate = with_data / len(pages) if pages else 0
        known = {path_key(p["url"]) for p in pages}
        extra = sum(1 for k in gsc if k not in known)
        print(f"Pages with Top GSC Queries: {with_data}/{len(pages)} ({rate:.0%}); GSC pages not in crawl: {extra}")
        if rate < 0.5:
            print("WARNING: under half the pages have queries. Check the property, the date range and the brand terms "
                  "before trusting the column.")

    qcol = "Top GSC Queries (impressions/clicks)"
    head = (["URL", "Review Needed", "Why Review", "Page Topic", "Current Page Title", "Current H1",
             "Suggested Page Title", "Suggested H1", "Suggested Keywords"] + ([qcol] if gsc is not None else []) + ["Notes"])
    width = {"URL": 45, "Review Needed": 11, "Why Review": 42, "Page Topic": 28, "Current Page Title": 40,
             "Current H1": 36, "Suggested Page Title": 40, "Suggested H1": 36, "Suggested Keywords": 38,
             qcol: 80, "Notes": 40}
    review = []
    for p in pages:
        row = [p["url"], "Yes" if p["flags"] else "No", "\n".join(t for _, t in p["flags"]), p["topic"],
               p["title"], p["h1"], None, None, ", ".join(p["keywords"])]
        if gsc is not None:
            row.append("; ".join(f"{q} ({i}/{c})" for q, i, c in p["queries"]) or NO_GSC_TEXT)
        row.append(" | ".join(p["notes"]))
        review.append(row)
    dhead = ["URL", "Page Path", "H1", "Title Tag"]
    data = [[p["url"], p["path"], p["h1"], p["title"]] for p in pages]

    wb = Workbook()
    ws = add_sheet(wb, True, "Review", head, review, [width[h] for h in head],
                   yellow=("Suggested Page Title", "Suggested H1"))
    ws.cell(1, head.index("Suggested Keywords") + 1).comment = Comment(
        "Target keywords chosen from the page's path, H1 and title, most important first. They are not search "
        "volume and are not validated against search data.", "Mavan")
    if gsc is not None:
        ws.cell(1, head.index(qcol) + 1).comment = Comment(
            f"Up to {TOP_QUERIES} non-branded queries per page by impressions (at least {MIN_IMPR} each), shown as "
            "impressions/clicks over the date range. For manual review; they are not ranked or judged.", "Mavan")
    for cell in ws[get_column_letter(2)][1:]:
        if cell.value == "Yes":
            cell.fill = PatternFill("solid", fgColor="F8CBAD")
    add_sheet(wb, False, "Data", dhead, data, [45, 36, 36, 40])
    os.makedirs(os.path.dirname(os.path.abspath(args.out)), exist_ok=True)
    wb.save(args.out)
    print(f"Saved {args.out}")

    by_flag = defaultdict(int)
    for p in pages:
        for code, _ in p["flags"]:
            by_flag[code] += 1
    print(f"Review Needed = Yes on {sum(1 for p in pages if p['flags'])} of {len(pages)} pages; flags: {dict(by_flag)}")

    checks = [("Every crawled URL once in Data and once in Review",
               len(data) == len(review) == len(crawl_keys) and {norm_url(r[0]) for r in review} == crawl_keys),
              ("Suggested Keywords: at most 5, no duplicates",
               all(len(p["keywords"]) <= 5 and len({dedupe_key(k) for k in p["keywords"]}) == len(p["keywords"]) for p in pages)),
              (f"Top GSC Queries: at most {TOP_QUERIES} per page, none under {MIN_IMPR} impressions, none branded",
               all(len(p["queries"]) <= TOP_QUERIES and all(i >= MIN_IMPR and not any(b in q for b in brand)
                                                              for q, i, _ in p["queries"]) for p in pages)),
              ("Every page has a topic", all(p["topic"] for p in pages)),
              ("Every page with no keywords is marked Review Needed = Yes", all(p["flags"] for p in pages if not p["keywords"]))]
    failed = False
    for name, passed in checks:
        print(("PASS  " if passed else "FAIL  ") + name)
        failed |= not passed
    sys.exit(1 if failed else 0)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    for name in ("inspect", "build"):
        sp = sub.add_parser(name)
        sp.add_argument("--crawl", required=True, help="Screaming Frog or Sitebulb export (CSV or xlsx)")
        sp.add_argument("--url-col")
        sp.add_argument("--h1-col")
        sp.add_argument("--title-col")
        sp.add_argument("--out", required=True, help="inspect: work folder. build: output .xlsx path")
        if name == "build":
            sp.add_argument("--keywords", required=True, help="work/keywords.tsv: the AI's topic and target keywords per page")
            sp.add_argument("--gsc", help="the Google Sheet of Search Console data (page, query, clicks, impressions), downloaded as .xlsx or the text export. Optional: without it there is no Top GSC Queries column")
            sp.add_argument("--gsc-tab", help="which tab to use when the file has several tabs with GSC columns")
            sp.add_argument("--brand", help="the client's brand terms, comma separated; queries containing one are dropped. Optional")
            sp.add_argument("--cities", help="the client's city or service-area cities, comma separated. Optional: without it the city check is skipped")
    args = ap.parse_args()
    cmd_inspect(args) if args.cmd == "inspect" else cmd_build(args)


if __name__ == "__main__":
    main()
